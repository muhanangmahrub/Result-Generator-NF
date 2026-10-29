"""Build the Result Test UAT Word document from parsed scenarios.

Document structure, taken from the reference file
``Result Test UAT AMF Segmentation.docx``:

* ``Heading 1`` - module title (sheet name in upper case by default);
* ``Heading 2`` - ``<Test Case ID><TAB><Test Case>`` per test case;
* the screenshots of that test case as inline pictures in ``Normal``
  paragraphs right underneath its heading, each exactly as wide as the
  printable area (6.27" for A4 with 1" margins), height scaled automatically;
* a test case without screenshots still gets its heading, plus the text
  ``[Belum ada screenshot]``.

Real Word heading styles are used, so the headings show up in the Navigation
Pane. Passing ``template`` reuses the user's own .docx for fonts, headers,
footers and logo: the document is opened and its body emptied, keeping the
section setup (and therefore the header/footer parts) intact.
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Union

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from .models import Module, ScenarioFile, TestCase

NO_SCREENSHOT_TEXT = "[Belum ada screenshot]"
MISSING_FILE_TEMPLATE = "[Screenshot tidak ditemukan: {name}]"
UNREADABLE_FILE_TEMPLATE = "[Screenshot tidak bisa dibaca: {name}]"
OUTPUT_PREFIX = "Result Test UAT"

#: Word may be localised; try the English style first, then Indonesian.
_HEADING_STYLES: Dict[int, Sequence[str]] = {
    1: ("Heading 1", "Judul 1"),
    2: ("Heading 2", "Judul 2"),
}
_HEADING_FALLBACK_SIZE = {1: Pt(16), 2: Pt(13)}

_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


# --------------------------------------------------------------------------- #
# Screenshot input
# --------------------------------------------------------------------------- #

@dataclass
class Screenshot:
    """One screenshot to insert, with an optional caption."""

    path: Path
    caption: str = ""

    @classmethod
    def coerce(cls, value: Any) -> "Screenshot":
        if isinstance(value, Screenshot):
            return value
        if isinstance(value, (str, Path)):
            return cls(path=Path(value))
        if isinstance(value, Mapping):
            return cls(path=Path(value["path"]), caption=str(value.get("caption", "")))
        raise TypeError(f"Screenshot tidak dikenali: {value!r}")


#: ``{TestCase.uid: [screenshot, ...]}`` - the mapping stored in project.json.
ScreenshotInput = Mapping[str, Sequence[Union[Screenshot, str, Path, Mapping[str, Any]]]]


def normalize_screenshots(screenshots: Optional[ScreenshotInput]) -> Dict[str, List[Screenshot]]:
    if not screenshots:
        return {}
    return {
        str(uid): [Screenshot.coerce(item) for item in items]
        for uid, items in screenshots.items()
    }


# --------------------------------------------------------------------------- #
# Reports
# --------------------------------------------------------------------------- #

@dataclass
class ModuleReport:
    """Per-module summary, shown to the user *before* exporting."""

    module: str
    title: str
    test_case_count: int = 0
    screenshot_count: int = 0
    without_screenshot: List[str] = field(default_factory=list)
    missing_files: List[str] = field(default_factory=list)

    @property
    def is_complete(self) -> bool:
        return not self.without_screenshot and not self.missing_files

    def to_dict(self) -> Dict[str, Any]:
        return {
            "module": self.module,
            "title": self.title,
            "test_case_count": self.test_case_count,
            "screenshot_count": self.screenshot_count,
            "without_screenshot": list(self.without_screenshot),
            "missing_files": list(self.missing_files),
        }


@dataclass
class ExportResult:
    """One written .docx plus the reports of the modules inside it."""

    path: Path
    modules: List[ModuleReport] = field(default_factory=list)

    @property
    def test_case_count(self) -> int:
        return sum(r.test_case_count for r in self.modules)

    @property
    def screenshot_count(self) -> int:
        return sum(r.screenshot_count for r in self.modules)

    @property
    def without_screenshot(self) -> List[str]:
        return [label for r in self.modules for label in r.without_screenshot]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "path": str(self.path),
            "modules": [r.to_dict() for r in self.modules],
        }


def summarize(
    scenario_file: ScenarioFile,
    screenshots: Optional[ScreenshotInput] = None,
    *,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
) -> List[ModuleReport]:
    """Report, per module, how many test cases still have no screenshot."""
    shots = normalize_screenshots(screenshots)
    naming = resolve_naming(scenario_file, modules=modules, titles=titles)
    reports: List[ModuleReport] = []
    for module in _selected_modules(scenario_file, modules):
        report = ModuleReport(
            module=module.name,
            title=naming[module.name].title,
        )
        for case in module.test_cases:
            report.test_case_count += 1
            items = shots.get(case.uid, [])
            if not items:
                report.without_screenshot.append(case.display_label)
                continue
            for shot in items:
                report.screenshot_count += 1
                if not shot.path.is_file():
                    report.missing_files.append(str(shot.path))
        reports.append(report)
    return reports


# --------------------------------------------------------------------------- #
# Module naming
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class ModuleNaming:
    """How one module is named in the document and in the output file name."""

    module: str
    title: str
    file_label: str
    source: str  # "override" | "module_name" | "project_name" | "sheet"


def _metadata_label(
    module: Module, merged: Mapping[str, str], single_module_workbook: bool
) -> tuple[str, str]:
    """Label for a module taken from metadata, as ``(label, source)``.

    Prefers the module's own metadata. The workbook-wide ``project_name`` is
    only borrowed when the whole workbook holds a single module - in a
    multi-module workbook it would name a module after a project name that
    belongs to a different sheet.
    """
    own = module.metadata
    if own.get("module_name"):
        return own["module_name"], "module_name"
    if own.get("project_name"):
        return own["project_name"], "project_name"
    if single_module_workbook and merged.get("project_name"):
        return merged["project_name"], "project_name"
    return "", "sheet"


def resolve_naming(
    scenario_file: ScenarioFile,
    *,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
) -> Dict[str, ModuleNaming]:
    """Work out the Heading 1 text and file label of every selected module.

    Order of preference:

    1. an explicit title from ``titles`` (what the user typed in the GUI) -
       used verbatim;
    2. the module's own ``module_name`` metadata;
    3. the module's own ``project_name`` metadata - this is what makes the real
       file produce "NEW VALIDATION AMF SEGMENTATION" instead of the sheet name
       "Validation AMF Segmentation";
    4. the workbook's ``project_name``, but only when the workbook has a single
       scenario module, so that a module never inherits another sheet's name;
    5. the sheet name.

    Automatic labels are upper-cased for Heading 1 and kept as written for the
    file name. If two modules would end up with the same automatic title, both
    fall back to their sheet name - a document with two identical Heading 1
    entries would be useless.
    """
    titles = titles or {}
    selected = _selected_modules(scenario_file, modules)
    merged = scenario_file.merged_metadata()
    single = len(scenario_file.modules) == 1

    naming: Dict[str, ModuleNaming] = {}
    for module in selected:
        if module.name in titles:
            text = titles[module.name]
            naming[module.name] = ModuleNaming(
                module=module.name, title=text, file_label=text, source="override"
            )
            continue
        label, source = _metadata_label(module, merged, single)
        if label:
            naming[module.name] = ModuleNaming(
                module=module.name,
                title=label.upper(),
                file_label=label,
                source=source,
            )
        else:
            naming[module.name] = ModuleNaming(
                module=module.name,
                title=module.title,
                file_label=module.name,
                source="sheet",
            )

    clashes = {
        title
        for title, count in Counter(item.title for item in naming.values()).items()
        if count > 1
    }
    for name, item in list(naming.items()):
        if item.source in ("override", "sheet") or item.title not in clashes:
            continue
        module = scenario_file.module(name)
        assert module is not None
        naming[name] = ModuleNaming(
            module=name, title=module.title, file_label=name, source="sheet"
        )
    return naming


def _selected_modules(
    scenario_file: ScenarioFile, modules: Optional[Iterable[str]]
) -> List[Module]:
    if modules is None:
        return list(scenario_file.modules)
    wanted = list(modules)
    unknown = [name for name in wanted if scenario_file.module(name) is None]
    if unknown:
        raise ValueError(
            "Module berikut tidak ada di hasil parse: " + ", ".join(unknown)
        )
    return [scenario_file.module(name) for name in wanted]  # type: ignore[misc]


# --------------------------------------------------------------------------- #
# Document primitives
# --------------------------------------------------------------------------- #

def _new_document(template: Optional[Path] = None):
    """Start an empty document, optionally keeping a template's look."""
    if template is None:
        return Document()
    template = Path(template)
    if not template.is_file():
        raise FileNotFoundError(f"File template tidak ditemukan: {template}")
    document = Document(str(template))
    _clear_body(document)
    _drop_orphan_images(document)
    return document


def _clear_body(document) -> None:
    """Remove every block of the body but keep the final section properties.

    ``w:sectPr`` holds page size, margins and the references to the header and
    footer parts, so dropping everything else leaves the user's page setup,
    logo and styles in place while removing their content.
    """
    body = document.element.body
    sect_pr_tag = qn("w:sectPr")
    for child in list(body):
        if child.tag == sect_pr_tag:
            continue
        body.remove(child)


def _drop_orphan_images(document) -> None:
    """Drop image parts the emptied body no longer references.

    Clearing the paragraphs does not remove the pictures from the .docx
    package, so using a finished result document as a style template would drag
    its old screenshots along - tens of megabytes of invisible payload. Only
    relationships of the document part are touched, so a logo in the header or
    footer (owned by the header/footer part) is kept.
    """
    part = document.part
    for rel_id, rel in list(part.rels.items()):
        if rel.reltype == RT.IMAGE:
            part.drop_rel(rel_id)


def _add_heading(document, text: str, level: int):
    for style_name in _HEADING_STYLES[level]:
        try:
            return document.add_paragraph(text, style=style_name)
        except KeyError:
            continue
    # No heading style at all in the template: fall back to bold text so the
    # document is still usable (it just will not show in the Navigation Pane).
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = _HEADING_FALLBACK_SIZE[level]
    return paragraph


def content_width(document) -> Emu:
    """Printable width of the last section - the width used for screenshots."""
    section = document.sections[-1]
    return Emu(section.page_width - section.left_margin - section.right_margin)


def _add_caption(document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def _add_screenshot(document, shot: Screenshot, width: Emu, report: ModuleReport) -> None:
    name = shot.path.name
    if not shot.path.is_file():
        document.add_paragraph(MISSING_FILE_TEMPLATE.format(name=name))
        return
    paragraph = document.add_paragraph()
    try:
        paragraph.add_run().add_picture(str(shot.path), width=width)
    except Exception:  # unsupported or corrupt image - keep the export going
        paragraph.add_run(UNREADABLE_FILE_TEMPLATE.format(name=name))
        if str(shot.path) not in report.missing_files:
            report.missing_files.append(str(shot.path))
        return
    if shot.caption:
        _add_caption(document, shot.caption)


def _add_test_case(document, case: TestCase, shots: Sequence[Screenshot],
                   width: Emu, report: ModuleReport) -> None:
    _add_heading(document, case.heading, 2)
    report.test_case_count += 1
    if not shots:
        document.add_paragraph(NO_SCREENSHOT_TEXT)
        report.without_screenshot.append(case.display_label)
        return
    for shot in shots:
        report.screenshot_count += 1
        if not shot.path.is_file():
            report.missing_files.append(str(shot.path))
        _add_screenshot(document, shot, width, report)


def _add_module(document, module: Module, shots: Dict[str, List[Screenshot]],
                title: str) -> ModuleReport:
    report = ModuleReport(module=module.name, title=title)
    _add_heading(document, title, 1)
    width = content_width(document)
    for case in module.test_cases:
        _add_test_case(document, case, shots.get(case.uid, []), width, report)
    return report


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #

def build_document(
    scenario_file: ScenarioFile,
    screenshots: Optional[ScreenshotInput] = None,
    *,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
    template: Optional[Path] = None,
    page_break_between_modules: bool = True,
):
    """Build one Document containing the selected modules.

    Returns ``(document, [ModuleReport, ...])``.
    """
    shots = normalize_screenshots(screenshots)
    naming = resolve_naming(scenario_file, modules=modules, titles=titles)
    document = _new_document(template)
    reports: List[ModuleReport] = []
    for index, module in enumerate(_selected_modules(scenario_file, modules)):
        if index and page_break_between_modules:
            document.add_page_break()
        reports.append(
            _add_module(document, module, shots, naming[module.name].title)
        )
    return document, reports


def safe_filename(name: str, *, fallback: str = "Output") -> str:
    """Strip characters Windows does not allow in a file name."""
    cleaned = _INVALID_FILENAME_CHARS.sub(" ", str(name))
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    return cleaned or fallback


def default_output_name(label: str) -> str:
    return f"{OUTPUT_PREFIX} {safe_filename(label)}.docx"


def export_per_module(
    scenario_file: ScenarioFile,
    screenshots: Optional[ScreenshotInput] = None,
    *,
    output_dir: Path,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
    template: Optional[Path] = None,
) -> List[ExportResult]:
    """One .docx per module: ``Result Test UAT <Module>.docx``."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    naming = resolve_naming(scenario_file, modules=modules, titles=titles)
    results: List[ExportResult] = []
    for module in _selected_modules(scenario_file, modules):
        # Pass the already resolved title so that splitting into one file per
        # module cannot rename a module (a single-module document would
        # otherwise be free to use the workbook-wide project name).
        document, reports = build_document(
            scenario_file,
            screenshots,
            modules=[module.name],
            titles={module.name: naming[module.name].title},
            template=template,
        )
        path = output_dir / default_output_name(naming[module.name].file_label)
        document.save(str(path))
        results.append(ExportResult(path=path, modules=reports))
    return results


def export_combined(
    scenario_file: ScenarioFile,
    screenshots: Optional[ScreenshotInput] = None,
    *,
    output_dir: Path,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
    template: Optional[Path] = None,
    project_name: Optional[str] = None,
    file_name: Optional[str] = None,
) -> ExportResult:
    """One .docx for every selected module: ``Result Test UAT <Project>.docx``."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    document, reports = build_document(
        scenario_file, screenshots, modules=modules, titles=titles, template=template
    )
    label = (
        project_name
        or scenario_file.merged_metadata().get("project_name")
        or scenario_file.path.stem
    )
    name = safe_filename(file_name) if file_name else default_output_name(label)
    if not name.lower().endswith(".docx"):
        name += ".docx"
    path = output_dir / name
    document.save(str(path))
    return ExportResult(path=path, modules=reports)


def export(
    scenario_file: ScenarioFile,
    screenshots: Optional[ScreenshotInput] = None,
    *,
    output_dir: Path,
    combined: bool = False,
    modules: Optional[Iterable[str]] = None,
    titles: Optional[Mapping[str, str]] = None,
    template: Optional[Path] = None,
    project_name: Optional[str] = None,
) -> List[ExportResult]:
    """Export per module (default) or as one combined document."""
    if combined:
        return [
            export_combined(
                scenario_file,
                screenshots,
                output_dir=output_dir,
                modules=modules,
                titles=titles,
                template=template,
                project_name=project_name,
            )
        ]
    return export_per_module(
        scenario_file,
        screenshots,
        output_dir=output_dir,
        modules=modules,
        titles=titles,
        template=template,
    )
