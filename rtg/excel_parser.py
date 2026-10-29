"""Read test scenarios from an Excel workbook.

Design rules enforced here:

* The workbook is opened **read-only in effect**: nothing is ever written back
  to the user's file.
* A sheet is a scenario sheet when a header row containing the mandatory
  columns (Test Case ID, Test Case, Test Step) can be found - never because of
  its name.
* Columns are located by matching *normalised* header text against the alias
  table from the config file. Column letters and positions are never used, and
  the header row may sit anywhere in the first ``header_scan_rows`` rows.
"""

from __future__ import annotations

import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .config import FieldConfig, load_config
from .errors import (
    CorruptWorkbookError,
    FileLockedError,
    MissingRequiredColumnsError,
    NoScenarioSheetError,
    ScenarioFileNotFoundError,
    SheetNotFoundError,
    UnsupportedFileFormatError,
)
from .models import STEP_FIELDS, Module, ScenarioFile, ScenarioGroup, TestCase, TestStep
from .normalize import cell_text, normalize_header, tokens_of

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}

Grid = Dict[Tuple[int, int], Any]


# --------------------------------------------------------------------------- #
# Manual mapping overrides (persisted in project.json)
# --------------------------------------------------------------------------- #

@dataclass
class SheetOverride:
    """Manual mapping decided by the user in the mapping dialog.

    ``columns`` maps a field name to a 1-based column index and wins over
    automatic detection; ``header_row`` forces the header row.
    """

    header_row: Optional[int] = None
    columns: Dict[str, int] = field(default_factory=dict)

    @classmethod
    def coerce(cls, value: Any) -> "SheetOverride":
        if value is None:
            return cls()
        if isinstance(value, SheetOverride):
            return value
        if isinstance(value, Mapping):
            columns = value.get("columns") or {}
            return cls(
                header_row=value.get("header_row"),
                columns={str(k): int(v) for k, v in columns.items()},
            )
        raise TypeError(f"Override sheet tidak dikenali: {value!r}")

    def to_dict(self) -> Dict[str, Any]:
        data: Dict[str, Any] = {"columns": dict(self.columns)}
        if self.header_row:
            data["header_row"] = self.header_row
        return data


def _coerce_overrides(overrides: Optional[Mapping[str, Any]]) -> Dict[str, SheetOverride]:
    if not overrides:
        return {}
    return {str(name): SheetOverride.coerce(value) for name, value in overrides.items()}


# --------------------------------------------------------------------------- #
# Inspection results
# --------------------------------------------------------------------------- #

@dataclass
class SheetInspection:
    """What we could work out about a single sheet, without parsing its rows."""

    name: str
    sheet_index: int
    is_scenario: bool = False
    header_row: Optional[int] = None
    column_map: Dict[str, int] = field(default_factory=dict)
    missing_required: List[str] = field(default_factory=list)
    headers: List[Tuple[int, str]] = field(default_factory=list)
    unmapped_columns: List[Tuple[int, str]] = field(default_factory=list)
    metadata: Dict[str, str] = field(default_factory=dict)
    data_row_estimate: int = 0
    required_fields: List[str] = field(default_factory=list)
    reason: str = ""

    @property
    def mapped_required(self) -> List[str]:
        return [f for f in self.required_fields if f in self.column_map]

    @property
    def is_mapping_candidate(self) -> bool:
        """Looks like a scenario table but needs the manual mapping dialog.

        Requires the Test Case ID column plus one other mandatory column.
        Test Case ID is the identity of a test case and its header wording is
        highly standardised, so a sheet without it is almost never a scenario
        table - whereas recap sheets such as "Main" or "Summary Scenario"
        routinely reuse the words "Test Case" and "Test Step" as headers for
        their count columns.

        This is only a hint for the sheet picker: the GUI still lets the user
        open the manual mapping dialog on any sheet.
        """
        if self.is_scenario or self.header_row is None:
            return False
        return "test_case_id" in self.column_map and len(self.mapped_required) >= 2

    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "sheet_index": self.sheet_index,
            "is_scenario": self.is_scenario,
            "header_row": self.header_row,
            "column_map": dict(self.column_map),
            "missing_required": list(self.missing_required),
            "headers": [list(item) for item in self.headers],
            "unmapped_columns": [list(item) for item in self.unmapped_columns],
            "metadata": dict(self.metadata),
            "data_row_estimate": self.data_row_estimate,
            "required_fields": list(self.required_fields),
            "reason": self.reason,
        }


@dataclass
class WorkbookInspection:
    """Per-sheet inspection of a whole workbook, for the sheet picker."""

    path: Path
    sheets: List[SheetInspection] = field(default_factory=list)

    @property
    def scenario_sheets(self) -> List[SheetInspection]:
        return [s for s in self.sheets if s.is_scenario]

    @property
    def mapping_candidates(self) -> List[SheetInspection]:
        return [s for s in self.sheets if s.is_mapping_candidate]

    @property
    def scenario_sheet_names(self) -> List[str]:
        return [s.name for s in self.scenario_sheets]

    def sheet(self, name: str) -> Optional[SheetInspection]:
        for item in self.sheets:
            if item.name == name:
                return item
        return None

    def merged_metadata(self) -> Dict[str, str]:
        merged: Dict[str, str] = {}
        for item in self.sheets:
            for key, value in item.metadata.items():
                if value and not merged.get(key):
                    merged[key] = value
        return merged

    def to_dict(self) -> Dict[str, Any]:
        return {"path": str(self.path), "sheets": [s.to_dict() for s in self.sheets]}


@dataclass
class _SheetScan:
    """Internal: inspection plus the resolved cell grid, reused when parsing."""

    inspection: SheetInspection
    grid: Grid
    max_row: int
    max_col: int


# --------------------------------------------------------------------------- #
# Workbook access
# --------------------------------------------------------------------------- #

def excel_lock_file(path: Path) -> Path:
    """Path of the ``~$name.xlsx`` lock file Excel creates for an open file."""
    path = Path(path)
    return path.with_name(f"~${path.name}")


def is_open_in_excel(path: Path) -> bool:
    """True when Excel appears to have the workbook open.

    Reading is still possible in that case; this is used to warn the user and
    to refuse writing to the file later on.
    """
    return excel_lock_file(path).exists()


def open_workbook(path: Path):
    """Open a workbook for reading, translating low-level errors."""
    path = Path(path)
    if not path.exists():
        raise ScenarioFileNotFoundError(
            f"File Excel tidak ditemukan: {path}"
        )
    if path.is_dir():
        raise UnsupportedFileFormatError(f"Path menunjuk ke folder, bukan file: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFileFormatError(
            f"Format '{suffix or 'tanpa ekstensi'}' tidak didukung. "
            "Simpan file sebagai .xlsx terlebih dahulu "
            f"(didukung: {', '.join(sorted(SUPPORTED_SUFFIXES))})."
        )
    try:
        with warnings.catch_warnings():
            # Scenario files often embed a logo as .emf, which openpyxl cannot
            # read and warns about. We never touch the workbook's images, so
            # the warning would only confuse the user.
            warnings.filterwarnings(
                "ignore", message=r".*cannot be read", category=UserWarning
            )
            # data_only=True -> formula cells return the value Excel last cached.
            return load_workbook(path, data_only=True, read_only=False, keep_links=False)
    except PermissionError as exc:
        raise FileLockedError(
            f"File '{path.name}' tidak bisa dibaca karena sedang dipakai proses lain. "
            "Tutup file di Excel lalu coba lagi.",
            detail=str(exc),
        ) from exc
    except zipfile.BadZipFile as exc:
        raise CorruptWorkbookError(
            f"File '{path.name}' tidak bisa dibuka sebagai workbook Excel. "
            "File mungkin rusak atau sebenarnya bukan .xlsx.",
            detail=str(exc),
        ) from exc
    except OSError as exc:  # pragma: no cover - platform dependent
        raise CorruptWorkbookError(
            f"Gagal membuka '{path.name}': {exc}", detail=str(exc)
        ) from exc


def _build_grid(ws: Worksheet) -> Tuple[Grid, int, int]:
    """Flatten a sheet into ``{(row, col): value}``, resolving merged ranges.

    openpyxl only stores a merged range's value in its top-left cell. Copying
    that value across the range makes both vertically merged Test Case ID cells
    and horizontally merged metadata labels behave like plain cells.
    """
    grid: Grid = {}
    max_row = 0
    max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            grid[(cell.row, cell.column)] = cell.value
            max_row = max(max_row, cell.row)
            max_col = max(max_col, cell.column)

    merged_cells = getattr(ws, "merged_cells", None)
    merged_ranges = list(merged_cells.ranges) if merged_cells is not None else []
    for merged in merged_ranges:
        value = grid.get((merged.min_row, merged.min_col))
        if value is None:
            continue
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                grid.setdefault((r, c), value)
                max_row = max(max_row, r)
                max_col = max(max_col, c)

    return grid, max_row, max_col


# --------------------------------------------------------------------------- #
# Column matching
# --------------------------------------------------------------------------- #

_EXACT = 1000
_EDGE = 700
_INSIDE = 500


def _contains_tokens(haystack: Sequence[str], needle: Sequence[str]) -> bool:
    n, m = len(haystack), len(needle)
    if m == 0 or m > n:
        return False
    for i in range(n - m + 1):
        if tuple(haystack[i:i + m]) == tuple(needle):
            return True
    return False


def _score_alias(header_tokens: Sequence[str], alias_tokens: Sequence[str]) -> int:
    """How well a normalised header matches one alias (0 = no match).

    Exact equality always beats a partial match, so a sheet containing both
    "Test Case" and "Test Case ID" maps each to the right field regardless of
    their order.
    """
    if tuple(header_tokens) == tuple(alias_tokens):
        return _EXACT + len(alias_tokens) * 10
    if not _contains_tokens(header_tokens, alias_tokens):
        return 0
    size = len(alias_tokens)
    at_edge = (
        tuple(header_tokens[:size]) == tuple(alias_tokens)
        or tuple(header_tokens[-size:]) == tuple(alias_tokens)
    )
    base = _EDGE if at_edge else _INSIDE
    extra = len(header_tokens) - size
    return base + size * 10 - extra


def match_columns(
    normalized_headers: Sequence[str],
    config: FieldConfig,
    *,
    first_column: int = 1,
) -> Dict[str, int]:
    """Map field names to 1-based column indexes for one header row.

    Greedy best-score assignment: each field and each column is used at most
    once, highest score first, ties broken by the left-most column so the
    result is deterministic.
    """
    scored: List[Tuple[int, int, str, int]] = []
    for offset, header in enumerate(normalized_headers):
        if not header:
            continue
        column = first_column + offset
        header_tokens = tokens_of(header)
        for field_name, aliases in config.aliases.items():
            best = 0
            for alias_tokens in aliases:
                score = _score_alias(header_tokens, alias_tokens)
                if score > best:
                    best = score
            if best:
                scored.append((-best, column, field_name, column))

    scored.sort()
    mapping: Dict[str, int] = {}
    used_columns: set = set()
    for _neg_score, _col_sort, field_name, column in scored:
        if field_name in mapping or column in used_columns:
            continue
        mapping[field_name] = column
        used_columns.add(column)
    return mapping


# --------------------------------------------------------------------------- #
# Header row detection
# --------------------------------------------------------------------------- #

@dataclass
class _HeaderCandidate:
    row: int
    mapping: Dict[str, int]
    raw: List[Tuple[int, str]]
    missing: List[str]


def _row_headers(grid: Grid, row: int, max_col: int) -> List[Tuple[int, str]]:
    return [(c, cell_text(grid.get((row, c)))) for c in range(1, max_col + 1)]


def _header_candidates(
    grid: Grid, max_row: int, max_col: int, config: FieldConfig
) -> List[_HeaderCandidate]:
    limit = min(max_row, max(1, config.header_scan_rows))
    candidates: List[_HeaderCandidate] = []
    for row in range(1, limit + 1):
        raw = _row_headers(grid, row, max_col)
        if sum(1 for _c, text in raw if text) < 2:
            continue  # a title or a single metadata label, not a header row
        normalized = [normalize_header(text) for _c, text in raw]
        mapping = match_columns(normalized, config)
        if not mapping:
            continue
        missing = [f for f in config.required_fields if f not in mapping]
        candidates.append(_HeaderCandidate(row=row, mapping=mapping, raw=raw, missing=missing))
    return candidates


def _pick_header(
    candidates: Sequence[_HeaderCandidate], forced_row: Optional[int]
) -> Optional[_HeaderCandidate]:
    if forced_row:
        for candidate in candidates:
            if candidate.row == forced_row:
                return candidate
        return None
    complete = [c for c in candidates if not c.missing]
    if complete:
        # Most recognised columns wins; the topmost row breaks ties.
        return min(complete, key=lambda c: (-len(c.mapping), c.row))
    if candidates:
        return min(candidates, key=lambda c: (-len(c.mapping), c.row))
    return None


# --------------------------------------------------------------------------- #
# Metadata above the header
# --------------------------------------------------------------------------- #

def _metadata_field(text: str, config: FieldConfig) -> Optional[str]:
    """Which metadata field a cell is a label for, or None."""
    normalized = normalize_header(text.rstrip(":").strip())
    if not normalized:
        return None
    tokens = tokens_of(normalized)
    best = 0
    found: Optional[str] = None
    for name, aliases in config.metadata_aliases.items():
        for alias_tokens in aliases:
            score = _score_alias(tokens, alias_tokens)
            if score > best:
                best, found = score, name
    return found


def extract_metadata(
    grid: Grid, header_row: int, max_col: int, config: FieldConfig
) -> Dict[str, str]:
    """Read ``label -> value`` metadata from the rows above the header.

    A label is any cell whose normalised text matches a metadata alias; the
    value is the first non-empty cell to its right on the same row. The scan
    stops at the next metadata label, because these templates put two
    label/value pairs side by side on one row ("Project Name | ... | Created By
    | ..."), so a label with an empty value must stay empty instead of
    borrowing its neighbour's value. The first occurrence of a field wins.
    """
    metadata: Dict[str, str] = {}
    if header_row <= 1:
        return metadata

    span = max(1, config.metadata_scan_columns)
    for row in range(1, header_row):
        for col in range(1, max_col + 1):
            label = cell_text(grid.get((row, col)))
            if not label:
                continue
            field_name = _metadata_field(label, config)
            if not field_name or field_name in metadata:
                continue
            value = ""
            for step in range(1, span + 1):
                candidate = cell_text(grid.get((row, col + step)))
                if not candidate:
                    continue
                if _metadata_field(candidate, config):
                    break  # reached the next label/value pair
                value = candidate
                break
            if value:
                metadata[field_name] = value
    return metadata


# --------------------------------------------------------------------------- #
# Sheet scanning
# --------------------------------------------------------------------------- #

def _scan_sheet(
    ws: Worksheet, sheet_index: int, config: FieldConfig, override: SheetOverride
) -> _SheetScan:
    grid, max_row, max_col = _build_grid(ws)
    inspection = SheetInspection(
        name=ws.title,
        sheet_index=sheet_index,
        required_fields=list(config.required_fields),
    )

    if not grid:
        inspection.reason = "Sheet kosong."
        return _SheetScan(inspection, grid, max_row, max_col)

    candidates = _header_candidates(grid, max_row, max_col, config)
    picked = _pick_header(candidates, override.header_row)

    if picked is None and override.header_row:
        row = override.header_row
        raw = _row_headers(grid, row, max_col)
        normalized = [normalize_header(text) for _c, text in raw]
        picked = _HeaderCandidate(
            row=row,
            mapping=match_columns(normalized, config),
            raw=raw,
            missing=[],
        )

    if picked is None:
        inspection.reason = (
            "Tidak ada baris header yang dikenali pada "
            f"{min(max_row, config.header_scan_rows)} baris pertama."
        )
        return _SheetScan(inspection, grid, max_row, max_col)

    mapping = dict(picked.mapping)
    for field_name, column in override.columns.items():
        mapping[field_name] = int(column)

    inspection.header_row = picked.row
    inspection.column_map = mapping
    inspection.headers = [(c, text) for c, text in picked.raw if text]
    mapped_columns = set(mapping.values())
    inspection.unmapped_columns = [
        (c, text) for c, text in picked.raw if text and c not in mapped_columns
    ]
    inspection.missing_required = [f for f in config.required_fields if f not in mapping]
    inspection.is_scenario = not inspection.missing_required
    inspection.metadata = extract_metadata(grid, picked.row, max_col, config)

    if "test_step" in mapping:
        step_col = mapping["test_step"]
        inspection.data_row_estimate = sum(
            1
            for r in range(picked.row + 1, max_row + 1)
            if cell_text(grid.get((r, step_col)))
        )
    if not inspection.is_scenario:
        inspection.reason = (
            "Kolom wajib tidak ditemukan: "
            + ", ".join(inspection.missing_required)
        )
    return _SheetScan(inspection, grid, max_row, max_col)


# --------------------------------------------------------------------------- #
# Row extraction
# --------------------------------------------------------------------------- #

class _FillDown:
    """Carries group columns downwards, resetting lower levels on key change."""

    def __init__(self, config: FieldConfig, mapping: Mapping[str, int]) -> None:
        self._groups = []
        for group in config.fill_down_groups:
            fields = tuple(f for f in group.fields if f in mapping)
            if not fields:
                continue
            key = group.key if group.key in mapping else fields[0]
            self._groups.append((key, fields))
        self._carry: Dict[str, str] = {}

    def apply(self, values: Dict[str, str]) -> Dict[str, str]:
        for index, (key, fields) in enumerate(self._groups):
            incoming = values.get(key, "")
            if incoming and incoming != self._carry.get(key):
                # New group starts here: drop stale values of this and deeper levels.
                for _key, deeper in self._groups[index:]:
                    for name in deeper:
                        self._carry.pop(name, None)
            for name in fields:
                value = values.get(name, "")
                if value:
                    self._carry[name] = value
                elif self._carry.get(name):
                    values[name] = self._carry[name]
        return values


def _looks_like_repeated_header(
    values: Mapping[str, str], mapping: Mapping[str, int], config: FieldConfig
) -> bool:
    """True when a data row is actually a repeat of the header row."""
    hits = 0
    for field_name in mapping:
        text = values.get(field_name, "")
        if not text:
            continue
        tokens = tokens_of(normalize_header(text))
        for alias_tokens in config.aliases.get(field_name, ()):  # type: ignore[arg-type]
            if tuple(tokens) == tuple(alias_tokens):
                hits += 1
                break
    return hits >= 2


def _read_rows(
    scan: _SheetScan, config: FieldConfig
) -> Tuple[List[Dict[str, Any]], List[str]]:
    inspection = scan.inspection
    mapping = inspection.column_map
    grid = scan.grid
    warnings: List[str] = []
    rows: List[Dict[str, Any]] = []

    fill_down = _FillDown(config, mapping)
    blank_run = 0
    row = (inspection.header_row or 0) + 1
    while row <= scan.max_row:
        values = {
            field_name: cell_text(grid.get((row, column)))
            for field_name, column in mapping.items()
        }
        if not any(values.values()):
            blank_run += 1
            if blank_run >= config.max_blank_rows:
                break
            row += 1
            continue
        blank_run = 0

        if _looks_like_repeated_header(values, mapping, config):
            row += 1
            continue

        values = fill_down.apply(values)
        if values.get("test_case_id") or values.get("test_case") or values.get("test_step"):
            values["_row"] = row
            rows.append(values)
        row += 1

    if blank_run >= config.max_blank_rows:
        first_blank = row - blank_run + 1
        remaining = [
            r
            for r in range(row + 1, scan.max_row + 1)
            if any(cell_text(grid.get((r, c))) for c in mapping.values())
        ]
        if remaining:
            warnings.append(
                f"Pembacaan dihentikan di baris {first_blank} setelah "
                f"{config.max_blank_rows} baris kosong berturut-turut; masih ada "
                f"data di baris {remaining[0]} yang tidak dibaca."
            )
    return rows, warnings


def _build_module(scan: _SheetScan, config: FieldConfig) -> Module:
    inspection = scan.inspection
    module = Module(
        name=inspection.name,
        sheet_index=inspection.sheet_index,
        header_row=inspection.header_row or 0,
        column_map=dict(inspection.column_map),
        metadata=dict(inspection.metadata),
        unmapped_columns=list(inspection.unmapped_columns),
    )

    rows, warnings = _read_rows(scan, config)
    module.warnings.extend(warnings)

    case_list: List[TestCase] = []
    occurrences: Dict[str, int] = {}
    occurrence_rows: Dict[str, List[int]] = {}
    ignored_rows: List[int] = []
    current: Optional[TestCase] = None

    for values in rows:
        test_case_id = values.get("test_case_id", "")
        name = values.get("test_case", "")
        if test_case_id:
            # A repeated Test Case ID is kept as a separate test case: it is
            # usually a typo in the workbook, not the same case continued. Rows
            # only continue the current case while the ID and the name agree.
            starts_new = (
                current is None
                or current.test_case_id != test_case_id
                or bool(
                    name and current.test_case
                    and name.strip() != current.test_case.strip()
                )
            )
        else:
            # No Test Case ID at all: keep rows of one unnamed case together.
            starts_new = current is None or bool(current.test_case_id)

        created = starts_new
        if starts_new:
            occurrence = 1
            if test_case_id:
                occurrence = occurrences.get(test_case_id, 0) + 1
                occurrences[test_case_id] = occurrence
                occurrence_rows.setdefault(test_case_id, []).append(values["_row"])
            current = TestCase(
                module=module.name,
                test_case_id=test_case_id,
                test_case=name,
                scenario_id=values.get("scenario_id", ""),
                test_scenario=values.get("test_scenario", ""),
                test_scenario_description=values.get("test_scenario_description", ""),
                row=values["_row"],
                occurrence=occurrence,
            )
            case_list.append(current)
        elif not current.test_case and name:
            current.test_case = name
        case = current

        if values.get("test_step"):
            step = TestStep(row=values["_row"])
            for name in STEP_FIELDS:
                if name in values:
                    setattr(step, name, values[name])
            case.steps.append(step)
        elif not created:
            # Neither a new test case nor a step: a stray row (for instance a
            # scenario name typed without any test case under it).
            ignored_rows.append(values["_row"])

    for duplicate, count in occurrences.items():
        if count < 2:
            continue
        found = occurrence_rows[duplicate]
        shown_rows = ", ".join(str(r) for r in found[:10])
        more_rows = f" dan {len(found) - 10} baris lain" if len(found) > 10 else ""
        module.warnings.append(
            f"Test Case ID '{duplicate}' muncul {count} kali di sheet "
            f"'{module.name}' (baris {shown_rows}{more_rows}); tetap dibuat sebagai "
            "test case terpisah. Cek apakah ID-nya salah ketik."
        )

    if ignored_rows:
        shown = ", ".join(str(r) for r in ignored_rows[:10])
        more = f" (dan {len(ignored_rows) - 10} baris lain)" if len(ignored_rows) > 10 else ""
        module.warnings.append(
            f"Sheet '{module.name}': baris {shown}{more} terisi sebagian tetapi "
            "tidak punya Test Step dan tidak membentuk test case baru, jadi "
            "diabaikan."
        )

    # Group test cases by scenario, preserving first-seen order.
    groups: Dict[Tuple[str, str], ScenarioGroup] = {}
    group_order: List[Tuple[str, str]] = []
    for case in case_list:
        group_key = (case.scenario_id, case.test_scenario)
        group = groups.get(group_key)
        if group is None:
            group = ScenarioGroup(
                scenario_id=case.scenario_id,
                test_scenario=case.test_scenario,
                test_scenario_description=case.test_scenario_description,
            )
            groups[group_key] = group
            group_order.append(group_key)
        group.test_cases.append(case)

    module.scenarios = [groups[key] for key in group_order]
    return module


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #

def inspect_workbook(
    path: Path,
    *,
    config: Optional[FieldConfig] = None,
    overrides: Optional[Mapping[str, Any]] = None,
) -> WorkbookInspection:
    """Inspect every sheet so the GUI can offer the scenario sheets as checkboxes."""
    config = config or load_config()
    override_map = _coerce_overrides(overrides)
    path = Path(path)
    workbook = open_workbook(path)
    try:
        inspection = WorkbookInspection(path=path)
        for index, name in enumerate(workbook.sheetnames):
            ws = workbook[name]
            scan = _scan_sheet(ws, index, config, override_map.get(name, SheetOverride()))
            inspection.sheets.append(scan.inspection)
        return inspection
    finally:
        workbook.close()


def parse_workbook(
    path: Path,
    sheets: Optional[Iterable[str]] = None,
    *,
    config: Optional[FieldConfig] = None,
    overrides: Optional[Mapping[str, Any]] = None,
    strict: bool = True,
) -> ScenarioFile:
    """Parse the selected sheets (default: every detected scenario sheet).

    Args:
        sheets: sheet names to parse. ``None`` parses all detected scenario
            sheets and raises :class:`NoScenarioSheetError` when there are none.
        overrides: ``{sheet_name: {"header_row": int, "columns": {...}}}`` from
            the manual mapping dialog, as stored in project.json.
        strict: when a sheet was named explicitly but lacks mandatory columns,
            raise :class:`MissingRequiredColumnsError` (so the GUI can open the
            mapping dialog) instead of skipping it.
    """
    config = config or load_config()
    override_map = _coerce_overrides(overrides)
    path = Path(path)
    workbook = open_workbook(path)
    try:
        available = list(workbook.sheetnames)
        requested = list(sheets) if sheets is not None else None
        if requested is not None:
            missing = [name for name in requested if name not in available]
            if missing:
                raise SheetNotFoundError(
                    "Sheet berikut tidak ada di file Excel: " + ", ".join(missing)
                    + ". Sheet yang tersedia: " + ", ".join(available) + "."
                )

        result = ScenarioFile(path=path)
        scans: List[_SheetScan] = []
        for index, name in enumerate(available):
            if requested is not None and name not in requested:
                continue
            ws = workbook[name]
            scans.append(
                _scan_sheet(ws, index, config, override_map.get(name, SheetOverride()))
            )

        usable = [scan for scan in scans if scan.inspection.is_scenario]
        if not usable:
            if requested is not None and strict:
                for scan in scans:
                    insp = scan.inspection
                    raise MissingRequiredColumnsError(
                        f"Sheet '{insp.name}' tidak punya kolom wajib: "
                        + ", ".join(insp.missing_required or list(config.required_fields))
                        + ". Gunakan mapping kolom manual untuk sheet ini.",
                        sheet_name=insp.name,
                        missing=insp.missing_required or list(config.required_fields),
                        headers=insp.headers,
                        header_row=insp.header_row,
                    )
            raise NoScenarioSheetError(
                "Tidak ada sheet skenario di file ini. Sheet skenario harus punya "
                "baris header yang memuat kolom "
                + ", ".join(config.required_fields)
                + ".",
                sheets=[scan.inspection for scan in scans],
            )

        if requested is not None and strict:
            for scan in scans:
                if not scan.inspection.is_scenario:
                    insp = scan.inspection
                    raise MissingRequiredColumnsError(
                        f"Sheet '{insp.name}' tidak punya kolom wajib: "
                        + ", ".join(insp.missing_required or list(config.required_fields))
                        + ". Gunakan mapping kolom manual untuk sheet ini.",
                        sheet_name=insp.name,
                        missing=insp.missing_required or list(config.required_fields),
                        headers=insp.headers,
                        header_row=insp.header_row,
                    )

        for scan in usable:
            module = _build_module(scan, config)
            if not module.test_cases:
                module.warnings.append(
                    f"Sheet '{module.name}' tidak punya baris data di bawah header "
                    f"(baris {module.header_row})."
                )
            result.modules.append(module)

        for module in result.modules:
            result.warnings.extend(module.warnings)

        if is_open_in_excel(path):
            result.warnings.append(
                f"File '{path.name}' sedang terbuka di Excel. Data tetap dibaca, "
                "tetapi pastikan perubahan terakhir sudah disimpan."
            )
        return result
    finally:
        workbook.close()
