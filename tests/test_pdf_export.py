"""Unit tests for the PDF export.

The actual conversion needs Microsoft Word or LibreOffice; those tests are
skipped when neither is installed. Everything else runs anywhere.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from rtg.pdf_export import (
    CONVERTERS,
    LIBREOFFICE,
    WORD,
    PdfExportError,
    available_converters,
    convert_many,
    default_converter,
    docx_to_pdf,
    find_soffice,
)


@pytest.fixture
def document(tmp_path):
    doc = docx.Document()
    doc.add_heading("MODULE CONTOH", 1)
    doc.add_paragraph("1.1\tTest case contoh", style="Heading 2")
    doc.add_paragraph("[Belum ada screenshot]")
    path = tmp_path / "Result Test UAT Contoh.docx"
    doc.save(str(path))
    return path


@pytest.fixture(scope="module")
def converter():
    name = default_converter()
    if name is None:                                  # pragma: no cover
        pytest.skip("tidak ada Word maupun LibreOffice di mesin ini")
    return name


# --------------------------------------------------------------------------- #
# Detection
# --------------------------------------------------------------------------- #

def test_available_converters_covers_both_names():
    infos = available_converters()
    assert {info.name for info in infos} == set(CONVERTERS)
    for info in infos:
        assert isinstance(info.available, bool)
        if not info.available:
            assert info.detail, f"{info.name} harus menjelaskan kenapa tidak tersedia"
    assert {info.label for info in infos} == {"Microsoft Word", "LibreOffice"}


def test_default_converter_is_one_of_the_available_ones():
    name = default_converter()
    if name is None:                                  # pragma: no cover
        assert all(not info.available for info in available_converters())
        return
    assert name in CONVERTERS
    assert next(i for i in available_converters() if i.name == name).available


def test_find_soffice_returns_a_file_or_nothing():
    found = find_soffice()
    assert found is None or Path(found).is_file()


# --------------------------------------------------------------------------- #
# Argument handling
# --------------------------------------------------------------------------- #

def test_missing_source(tmp_path):
    with pytest.raises(PdfExportError, match="tidak ditemukan"):
        docx_to_pdf(tmp_path / "ghost.docx")


def test_wrong_file_type(tmp_path):
    other = tmp_path / "catatan.txt"
    other.write_text("bukan word", encoding="utf-8")
    with pytest.raises(PdfExportError, match="Hanya file Word"):
        docx_to_pdf(other)


def test_unknown_converter(document):
    with pytest.raises(PdfExportError, match="tidak dikenali"):
        docx_to_pdf(document, converter="telepati")


def test_existing_pdf_is_refused_when_overwrite_is_off(document, tmp_path):
    target = tmp_path / "ada.pdf"
    target.write_bytes(b"%PDF-1.7 lama")
    with pytest.raises(PdfExportError, match="sudah ada"):
        docx_to_pdf(document, target, overwrite=False)


def test_libreoffice_missing_is_explained(document, tmp_path, monkeypatch):
    monkeypatch.setattr("rtg.pdf_export.find_soffice", lambda: None)
    with pytest.raises(PdfExportError, match="LibreOffice"):
        docx_to_pdf(document, tmp_path / "x.pdf", converter=LIBREOFFICE)


def test_no_converter_at_all_is_explained(document, tmp_path, monkeypatch):
    monkeypatch.setattr("rtg.pdf_export.default_converter", lambda: None)
    with pytest.raises(PdfExportError, match="Tidak ada converter PDF"):
        docx_to_pdf(document, tmp_path / "x.pdf")


# --------------------------------------------------------------------------- #
# Real conversion
# --------------------------------------------------------------------------- #

def test_convert_a_document(document, converter):
    pdf = docx_to_pdf(document, converter=converter)

    assert pdf == document.with_suffix(".pdf")
    assert pdf.is_file()
    assert pdf.read_bytes().startswith(b"%PDF")
    assert pdf.stat().st_size > 500


def test_convert_to_a_chosen_path(document, converter, tmp_path):
    target = tmp_path / "pdf" / "Hasil UAT.pdf"
    pdf = docx_to_pdf(document, target, converter=converter)
    assert pdf == target
    assert target.read_bytes().startswith(b"%PDF")


def test_overwriting_an_existing_pdf(document, converter, tmp_path):
    target = tmp_path / "ada.pdf"
    target.write_bytes(b"isi lama")
    pdf = docx_to_pdf(document, target, converter=converter)
    assert pdf.read_bytes().startswith(b"%PDF")


def test_convert_many(document, converter, tmp_path):
    second = tmp_path / "Result Test UAT Kedua.docx"
    doc = docx.Document()
    doc.add_heading("MODULE KEDUA", 1)
    doc.save(str(second))

    out = tmp_path / "pdf"
    produced = convert_many([document, second], converter=converter, output_dir=out)

    assert [p.name for p in produced] == [
        "Result Test UAT Contoh.pdf", "Result Test UAT Kedua.pdf",
    ]
    assert all(p.read_bytes().startswith(b"%PDF") for p in produced)
