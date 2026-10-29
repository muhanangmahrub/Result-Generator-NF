"""Unit tests for the project folder layer."""

from __future__ import annotations

import json

import docx
import pytest

from rtg.config import load_config
from rtg.project import (
    PROJECT_FILENAME,
    SCHEMA_VERSION,
    Project,
    ProjectError,
    ScreenshotEntry,
)
from tests.fixtures import build_workbook, make_png, simple_workbook, write_scenario_sheet

CONFIG = load_config()

ROWS_TWO_CASES = [
    {"scenario_id": "SC-01", "test_scenario": "Alpha", "test_case_id": "1.1",
     "test_case": "Case Satu", "test_step": "step a"},
    {"test_step": "step b"},
    {"test_case_id": "1.2", "test_case": "Case Dua", "test_step": "step c"},
]
ROWS_ONE_CASE = [
    {"scenario_id": "SC-01", "test_scenario": "Alpha", "test_case_id": "1.1",
     "test_case": "Case Satu", "test_step": "step a"},
]


@pytest.fixture
def excel(tmp_path):
    return simple_workbook(tmp_path / "src" / "Scenario.xlsx", rows=ROWS_TWO_CASES)


@pytest.fixture
def project(tmp_path, excel):
    return Project.create(tmp_path / "proj", excel, config=CONFIG)


@pytest.fixture
def shot(tmp_path):
    return make_png(tmp_path / "incoming" / "capture.png", 1920, 1080)


# --------------------------------------------------------------------------- #
# Create / open / save
# --------------------------------------------------------------------------- #

def test_create_builds_the_folder_layout(tmp_path, excel):
    project = Project.create(tmp_path / "proj", excel, config=CONFIG)

    assert project.project_file.is_file()
    assert project.screenshots_dir.is_dir()
    assert project.output_dir.is_dir()
    assert project.name == "proj"
    assert project.selected_sheets == ["Scenario"]
    assert project.known_uids == ["Scenario::1.1", "Scenario::1.2"]
    assert project.hotkey == "Ctrl+Shift+S"


def test_project_json_contents(project):
    data = json.loads(project.project_file.read_text(encoding="utf-8"))

    assert data["schema_version"] == SCHEMA_VERSION
    assert data["selected_sheets"] == ["Scenario"]
    assert data["output"] == {"mode": "per_module", "template": None}
    assert data["created_at"] and data["updated_at"]
    assert data["screenshots"] == {}


def test_create_refuses_an_existing_project(tmp_path, excel):
    Project.create(tmp_path / "proj", excel, config=CONFIG)
    with pytest.raises(ProjectError, match="sudah berisi project"):
        Project.create(tmp_path / "proj", excel, config=CONFIG)
    # ...unless explicitly allowed
    assert Project.create(tmp_path / "proj", excel, config=CONFIG, exist_ok=True)


def test_create_accepts_capture_settings(tmp_path, excel):
    project = Project.create(
        tmp_path / "proj", excel, config=CONFIG,
        hotkey="Ctrl+Shift+F", capture_mode="region", capture_monitor=2,
    )
    assert project.hotkey == "Ctrl+Shift+F"
    assert project.capture_mode == "region"
    assert project.capture_monitor == 2

    reopened = Project.open(project.folder, config=CONFIG)
    assert reopened.hotkey == "Ctrl+Shift+F"
    assert reopened.capture_mode == "region"


def test_create_falls_back_to_the_defaults(tmp_path, excel):
    project = Project.create(tmp_path / "proj", excel, config=CONFIG)
    assert project.hotkey == "Ctrl+Shift+S"
    assert project.capture_mode == "fullscreen"


def test_create_ignores_an_unknown_capture_mode(tmp_path, excel):
    project = Project.create(
        tmp_path / "proj", excel, config=CONFIG, capture_mode="telepati"
    )
    assert project.capture_mode == "fullscreen"


def test_create_with_missing_excel(tmp_path):
    with pytest.raises(ProjectError, match="tidak ditemukan"):
        Project.create(tmp_path / "proj", tmp_path / "nope.xlsx", config=CONFIG)


def test_create_with_explicit_sheet_selection(tmp_path):
    def builder(wb):
        write_scenario_sheet(wb.create_sheet("Module A"))
        write_scenario_sheet(wb.create_sheet("Module B"))

    path = build_workbook(tmp_path / "two.xlsx", builder)
    project = Project.create(tmp_path / "p", path, sheets=["Module B"], config=CONFIG)
    assert project.selected_sheets == ["Module B"]
    assert all(uid.startswith("Module B::") for uid in project.known_uids)


def test_open_round_trip(project, shot):
    project.add_screenshot("Scenario::1.1", shot, caption="Bukti 1")
    project.module_titles["Scenario"] = "JUDUL SAYA"
    project.metadata["project_name"] = "Proyek X"
    project.output_mode = "combined"
    project.hotkey = "Ctrl+Alt+S"
    project.save()

    reopened = Project.open(project.folder, config=CONFIG)

    assert reopened.excel_path == project.excel_path
    assert reopened.module_titles == {"Scenario": "JUDUL SAYA"}
    assert reopened.metadata["project_name"] == "Proyek X"
    assert reopened.output_mode == "combined"
    assert reopened.hotkey == "Ctrl+Alt+S"
    assert reopened.screenshot_count("Scenario::1.1") == 1
    assert reopened.entries("Scenario::1.1")[0].caption == "Bukti 1"
    assert reopened.known_uids == project.known_uids


def test_open_accepts_the_project_file_itself(project):
    assert Project.open(project.project_file, config=CONFIG).folder == project.folder


def test_open_non_project_folder(tmp_path):
    with pytest.raises(ProjectError, match="bukan folder project"):
        Project.open(tmp_path, config=CONFIG)


def test_open_corrupt_project_json(tmp_path, excel):
    project = Project.create(tmp_path / "proj", excel, config=CONFIG)
    project.project_file.write_text("{ bukan json", encoding="utf-8")
    with pytest.raises(ProjectError, match="rusak"):
        Project.open(project.folder, config=CONFIG)


def test_open_rejects_a_newer_schema(project):
    data = json.loads(project.project_file.read_text(encoding="utf-8"))
    data["schema_version"] = SCHEMA_VERSION + 1
    project.project_file.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ProjectError, match="versi aplikasi yang lebih baru"):
        Project.open(project.folder, config=CONFIG)


def test_save_leaves_no_temp_file(project):
    project.save()
    assert list(project.folder.glob("*.tmp")) == []


def test_excel_inside_the_project_folder_is_stored_relative(tmp_path):
    folder = tmp_path / "proj"
    folder.mkdir()
    excel = simple_workbook(folder / "Scenario.xlsx", rows=ROWS_TWO_CASES)
    project = Project.create(folder, excel, config=CONFIG, exist_ok=True)

    data = json.loads(project.project_file.read_text(encoding="utf-8"))
    assert data["excel_path"] == "Scenario.xlsx"
    assert Project.open(folder, config=CONFIG).known_uids == project.known_uids


def test_metadata_is_captured_on_create(tmp_path):
    from tests.fixtures import DEFAULT_METADATA

    excel = simple_workbook(
        tmp_path / "meta.xlsx", header_row=9, metadata=DEFAULT_METADATA,
        metadata_start_row=2, rows=ROWS_TWO_CASES,
    )
    project = Project.create(tmp_path / "p", excel, config=CONFIG)
    assert project.metadata["project_name"] == "New Validation AMF"


def test_effective_metadata_prefers_user_edits(project):
    project.metadata["project_name"] = "Diedit Manual"
    assert project.effective_metadata()["project_name"] == "Diedit Manual"


# --------------------------------------------------------------------------- #
# Screenshot mapping
# --------------------------------------------------------------------------- #

def test_add_screenshot_path_layout(project, shot):
    entry = project.add_screenshot("Scenario::1.1", shot, caption="Langkah 1")

    parts = entry.file.split("/")
    assert parts[0] == "screenshots"
    assert parts[1] == "scenario"                      # <module-slug>
    assert parts[2] == "1.1_case-satu"                 # <TCID>_<slug>
    assert parts[3].startswith("1.1_001_")             # <TCID>_<urutan>_<ts>
    assert parts[3].endswith(".png")
    assert project.absolute(entry).is_file()
    assert entry.caption == "Langkah 1"
    assert entry.captured_at


def test_add_screenshot_copies_by_default_and_can_move(project, shot, tmp_path):
    project.add_screenshot("Scenario::1.1", shot)
    assert shot.is_file()                              # source kept

    second = make_png(tmp_path / "incoming" / "second.png")
    project.add_screenshot("Scenario::1.1", second, move=True)
    assert not second.exists()                         # source moved away
    assert project.screenshot_count("Scenario::1.1") == 2


def test_sequence_numbers_increase(project, tmp_path):
    for index in range(1, 4):
        png = make_png(tmp_path / "in" / f"{index}.png")
        entry = project.add_screenshot("Scenario::1.1", png)
        assert f"_{index:03d}_" in entry.file


def test_add_screenshot_to_unknown_test_case(project, shot):
    with pytest.raises(ProjectError, match="tidak ada di skenario"):
        project.add_screenshot("Scenario::9.9", shot)


def test_add_screenshot_with_missing_image(project, tmp_path):
    with pytest.raises(ProjectError, match="File gambar tidak ditemukan"):
        project.add_screenshot("Scenario::1.1", tmp_path / "ghost.png")


def test_same_test_case_id_in_two_modules_uses_separate_folders(tmp_path):
    def builder(wb):
        write_scenario_sheet(wb.create_sheet("Module A"), rows=ROWS_ONE_CASE)
        write_scenario_sheet(wb.create_sheet("Module B"), rows=ROWS_ONE_CASE)

    excel = build_workbook(tmp_path / "dup.xlsx", builder)
    project = Project.create(tmp_path / "p", excel, config=CONFIG)

    a = project.add_screenshot("Module A::1.1", make_png(tmp_path / "a.png"))
    b = project.add_screenshot("Module B::1.1", make_png(tmp_path / "b.png"))

    assert "module-a" in a.file and "module-b" in b.file
    assert project.absolute(a) != project.absolute(b)
    assert project.screenshot_count("Module A::1.1") == 1
    assert project.screenshot_count("Module B::1.1") == 1


def test_set_caption(project, shot):
    project.add_screenshot("Scenario::1.1", shot)
    project.set_caption("Scenario::1.1", 0, "Caption baru")
    assert Project.open(project.folder, config=CONFIG).entries("Scenario::1.1")[0].caption == "Caption baru"


def test_reorder_screenshot(project, tmp_path):
    files = [
        project.add_screenshot("Scenario::1.1", make_png(tmp_path / "in" / f"{i}.png")).file
        for i in range(3)
    ]
    project.reorder_screenshot("Scenario::1.1", 2, 0)
    assert [e.file for e in project.entries("Scenario::1.1")] == [files[2], files[0], files[1]]


def test_reorder_out_of_range(project, shot):
    project.add_screenshot("Scenario::1.1", shot)
    with pytest.raises(ProjectError, match="di luar jangkauan"):
        project.reorder_screenshot("Scenario::1.1", 5, 0)


def test_move_screenshot_to_another_test_case(project, shot):
    entry = project.add_screenshot("Scenario::1.1", shot)
    old_path = project.absolute(entry)

    moved = project.move_screenshot("Scenario::1.1", 0, "Scenario::1.2")

    assert project.screenshot_count("Scenario::1.1") == 0
    assert project.screenshot_count("Scenario::1.2") == 1
    assert "1.2_case-dua" in moved.file
    assert not old_path.exists()
    assert project.absolute(moved).is_file()


def test_move_screenshot_to_unknown_test_case(project, shot):
    project.add_screenshot("Scenario::1.1", shot)
    with pytest.raises(ProjectError, match="tujuan tidak ada"):
        project.move_screenshot("Scenario::1.1", 0, "Scenario::9.9")


def test_delete_screenshot(project, shot):
    entry = project.add_screenshot("Scenario::1.1", shot)
    path = project.absolute(entry)

    project.delete_screenshot("Scenario::1.1", 0)
    assert project.screenshot_count("Scenario::1.1") == 0
    assert not path.exists()


def test_delete_screenshot_keeping_the_file(project, shot):
    entry = project.add_screenshot("Scenario::1.1", shot)
    path = project.absolute(entry)
    project.delete_screenshot("Scenario::1.1", 0, delete_file=False)
    assert path.is_file()


def test_missing_files_are_detected(project, shot):
    entry = project.add_screenshot("Scenario::1.1", shot)
    project.absolute(entry).unlink()
    missing = project.missing_files()
    assert [uid for uid, _entry in missing] == ["Scenario::1.1"]


# --------------------------------------------------------------------------- #
# Re-import
# --------------------------------------------------------------------------- #

def test_reimport_without_changes(project, shot):
    project.add_screenshot("Scenario::1.1", shot)
    report = project.reimport()

    assert report.has_changes is False
    assert report.summary_lines() == ["Tidak ada perubahan skenario."]
    assert project.screenshot_count("Scenario::1.1") == 1


def test_reimport_keeps_screenshots_and_orphans_the_lost_ones(project, shot, excel):
    project.add_screenshot("Scenario::1.1", shot)
    entry = project.add_screenshot("Scenario::1.2", shot, caption="Punya 1.2")
    kept_file = project.absolute(entry)

    # The Excel file now only describes test case 1.1.
    simple_workbook(excel, rows=ROWS_ONE_CASE)
    report = project.reimport()

    assert report.removed_test_cases == ["Scenario::1.2"]
    assert report.orphaned == {"Scenario::1.2": 1}
    assert report.orphaned_count == 1
    assert project.screenshot_count("Scenario::1.1") == 1     # untouched
    assert project.screenshot_count("Scenario::1.2") == 0
    assert project.orphan_uids == ["Scenario::1.2"]
    # The file itself is never deleted.
    assert kept_file.is_file()
    assert "orphan" in " ".join(report.summary_lines())


def test_orphans_are_restored_when_the_test_case_comes_back(project, shot, excel):
    project.add_screenshot("Scenario::1.2", shot, caption="Punya 1.2")

    simple_workbook(excel, rows=ROWS_ONE_CASE)
    project.reimport()
    assert project.orphan_count() == 1

    simple_workbook(excel, rows=ROWS_TWO_CASES)
    report = project.reimport()

    assert report.restored == {"Scenario::1.2": 1}
    assert project.orphan_count() == 0
    assert project.screenshot_count("Scenario::1.2") == 1
    assert project.entries("Scenario::1.2")[0].caption == "Punya 1.2"


def test_orphans_survive_save_and_open(project, shot, excel):
    project.add_screenshot("Scenario::1.2", shot)
    simple_workbook(excel, rows=ROWS_ONE_CASE)
    project.reimport()

    reopened = Project.open(project.folder, config=CONFIG)
    assert reopened.orphan_uids == ["Scenario::1.2"]
    assert reopened.orphan_count() == 1


def test_restore_orphan_to_another_test_case(project, shot, excel):
    project.add_screenshot("Scenario::1.2", shot, caption="Pindah")
    simple_workbook(excel, rows=ROWS_ONE_CASE)
    project.reimport()

    project.restore_orphan("Scenario::1.2", "Scenario::1.1")
    assert project.orphan_count() == 0
    assert project.screenshot_count("Scenario::1.1") == 1
    assert project.entries("Scenario::1.1")[0].caption == "Pindah"


def test_discard_orphans(project, shot, excel):
    entry = project.add_screenshot("Scenario::1.2", shot)
    path = project.absolute(entry)
    simple_workbook(excel, rows=ROWS_ONE_CASE)
    project.reimport()

    assert project.discard_orphans("Scenario::1.2", delete_files=True) == 1
    assert project.orphan_count() == 0
    assert not path.exists()


def test_restore_orphan_errors(project):
    with pytest.raises(ProjectError, match="Tidak ada screenshot orphan"):
        project.restore_orphan("Scenario::9.9", "Scenario::1.1")


def test_reimport_detects_new_test_cases(project, excel):
    rows = ROWS_TWO_CASES + [
        {"test_case_id": "1.3", "test_case": "Case Tiga", "test_step": "step d"}
    ]
    simple_workbook(excel, rows=rows)
    report = project.reimport()

    assert report.added_test_cases == ["Scenario::1.3"]
    assert "Scenario::1.3" in project.known_uids


def test_reimport_can_change_the_sheet_selection(tmp_path):
    def builder(wb):
        write_scenario_sheet(wb.create_sheet("Module A"), rows=ROWS_ONE_CASE)
        write_scenario_sheet(wb.create_sheet("Module B"), rows=ROWS_ONE_CASE)

    excel = build_workbook(tmp_path / "two.xlsx", builder)
    project = Project.create(tmp_path / "p", excel, sheets=["Module A"], config=CONFIG)
    project.module_titles["Module A"] = "JUDUL A"

    report = project.reimport(sheets=["Module B"])

    assert report.added_modules == ["Module B"]
    assert report.removed_modules == ["Module A"]
    assert project.selected_sheets == ["Module B"]
    # The title of a module that is gone is dropped too.
    assert "Module A" not in project.module_titles


def test_reimport_can_point_at_another_excel_file(project, tmp_path):
    other = simple_workbook(tmp_path / "other" / "Lain.xlsx", rows=ROWS_ONE_CASE)
    project.reimport(excel_path=other)
    assert project.excel_path == other
    assert project.known_uids == ["Scenario::1.1"]


def test_reimport_with_missing_excel(project, tmp_path):
    with pytest.raises(ProjectError, match="tidak ditemukan"):
        project.reimport(excel_path=tmp_path / "ghost.xlsx")


def test_reimport_persists_immediately(project, shot, excel):
    project.add_screenshot("Scenario::1.2", shot)
    simple_workbook(excel, rows=ROWS_ONE_CASE)
    project.reimport()

    data = json.loads(project.project_file.read_text(encoding="utf-8"))
    assert "Scenario::1.2" in data["orphan_screenshots"]
    assert "Scenario::1.2" not in data.get("screenshots", {})


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #

def test_export_screenshots_shape(project, shot):
    project.add_screenshot("Scenario::1.1", shot, caption="Satu")
    mapping = project.export_screenshots()

    assert list(mapping) == ["Scenario::1.1"]
    item = mapping["Scenario::1.1"][0]
    assert item.path.is_file()
    assert item.caption == "Satu"


def test_summarize_through_the_project(project, shot):
    project.add_screenshot("Scenario::1.1", shot)
    reports = project.summarize()

    assert reports[0].title == "SCENARIO"
    assert reports[0].test_case_count == 2
    assert reports[0].screenshot_count == 1
    assert reports[0].without_screenshot == ["1.2  Case Dua"]


def test_generate_writes_into_the_output_folder(project, shot):
    project.add_screenshot("Scenario::1.1", shot, caption="Bukti")
    results = project.generate()

    assert len(results) == 1
    produced = results[0].path
    assert produced.parent == project.output_dir
    assert produced.name == "Result Test UAT Scenario.docx"

    document = docx.Document(str(produced))
    assert len(document.inline_shapes) == 1
    assert "Bukti" in [p.text for p in document.paragraphs]


def test_generate_combined_mode(tmp_path):
    def builder(wb):
        write_scenario_sheet(wb.create_sheet("Module A"), rows=ROWS_ONE_CASE)
        write_scenario_sheet(wb.create_sheet("Module B"), rows=ROWS_ONE_CASE)

    excel = build_workbook(tmp_path / "two.xlsx", builder)
    project = Project.create(tmp_path / "p", excel, config=CONFIG)
    project.output_mode = "combined"
    project.metadata["project_name"] = "Proyek Gabungan"

    results = project.generate()
    assert len(results) == 1
    assert results[0].path.name == "Result Test UAT Proyek Gabungan.docx"
    assert len(results[0].modules) == 2


def test_generate_uses_module_titles(project):
    project.module_titles["Scenario"] = "JUDUL DIEDIT"
    results = project.generate()
    document = docx.Document(str(results[0].path))
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
        "JUDUL DIEDIT"
    ]


def test_generate_with_a_missing_template(project, tmp_path):
    project.template = str(tmp_path / "ghost.docx")
    with pytest.raises(ProjectError, match="template tidak ditemukan"):
        project.generate()


def test_generate_with_a_template(project, tmp_path):
    template = tmp_path / "template.docx"
    doc = docx.Document()
    doc.add_paragraph("isi lama")
    doc.save(str(template))
    project.template = str(template)

    results = project.generate()
    texts = [p.text for p in docx.Document(str(results[0].path)).paragraphs]
    assert "isi lama" not in texts


def test_screenshot_entry_accepts_a_bare_string():
    assert ScreenshotEntry.from_dict("screenshots/a.png").file == "screenshots/a.png"
