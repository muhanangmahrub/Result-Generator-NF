"""Unit tests for the Word result document builder."""

from __future__ import annotations

import docx
import pytest
from docx.shared import Inches, Pt

from rtg.config import load_config
from rtg.docx_builder import (
    NO_SCREENSHOT_TEXT,
    Screenshot,
    build_document,
    content_width,
    default_output_name,
    export,
    export_combined,
    export_per_module,
    resolve_naming,
    safe_filename,
    summarize,
)
from rtg.excel_parser import parse_workbook
from tests.fixtures import SAMPLE_ROWS, build_workbook, make_png, simple_workbook, write_scenario_sheet

CONFIG = load_config()


@pytest.fixture
def scenario(tmp_path):
    """One module ('Scenario') with test cases 1.1 (2 steps) and 1.2 (1 step)."""
    path = simple_workbook(tmp_path / "scenario.xlsx")
    return parse_workbook(path, config=CONFIG)


@pytest.fixture
def two_modules(tmp_path):
    def builder(wb):
        write_scenario_sheet(wb.create_sheet("AMF Segmentation"))
        write_scenario_sheet(wb.create_sheet("Billing Adjustment"))

    path = build_workbook(tmp_path / "two.xlsx", builder)
    return parse_workbook(path, config=CONFIG)


def _texts(document, style=None):
    return [
        p.text
        for p in document.paragraphs
        if style is None or p.style.name == style
    ]


def _image_count(document):
    return len(document.inline_shapes)


# --------------------------------------------------------------------------- #
# Structure
# --------------------------------------------------------------------------- #

def test_heading_structure(scenario):
    document, reports = build_document(scenario)

    assert _texts(document, "Heading 1") == ["SCENARIO"]
    assert _texts(document, "Heading 2") == [
        "1.1\tMove Partner Onboarding File via Sharefolder",
        "1.2\tValidasi File Hasil Onboarding",
    ]
    assert reports[0].test_case_count == 2


def test_real_heading_styles_are_used(scenario):
    """Heading 1/2 must be genuine Word styles so the Navigation Pane works."""
    document, _ = build_document(scenario)
    styles = [p.style.name for p in document.paragraphs if p.text.strip()]
    assert "Heading 1" in styles
    assert "Heading 2" in styles


def test_test_case_without_screenshot_still_gets_a_heading(scenario):
    document, reports = build_document(scenario)

    assert _texts(document, "Normal").count(NO_SCREENSHOT_TEXT) == 2
    assert reports[0].without_screenshot == [
        "1.1  Move Partner Onboarding File via Sharefolder",
        "1.2  Validasi File Hasil Onboarding",
    ]
    assert reports[0].is_complete is False


def test_screenshots_are_inserted_in_order_under_their_heading(scenario, tmp_path):
    first = make_png(tmp_path / "shots" / "a.png")
    second = make_png(tmp_path / "shots" / "b.png")
    screenshots = {"Scenario::1.1": [first, second]}

    document, reports = build_document(scenario, screenshots)

    assert _image_count(document) == 2
    assert reports[0].screenshot_count == 2
    # 1.1 has pictures, 1.2 still carries the placeholder.
    assert reports[0].without_screenshot == ["1.2  Validasi File Hasil Onboarding"]

    # Order of blocks: H1, H2(1.1), img, img, H2(1.2), placeholder
    blocks = [
        (p.style.name, "IMG" if "graphicData" in p._p.xml else p.text)
        for p in document.paragraphs
    ]
    assert blocks[0][0] == "Heading 1"
    assert blocks[1][1].startswith("1.1")
    assert blocks[2][1] == "IMG"
    assert blocks[3][1] == "IMG"
    assert blocks[4][1].startswith("1.2")
    assert blocks[5][1] == NO_SCREENSHOT_TEXT


def test_screenshot_width_matches_printable_area(scenario, tmp_path):
    png = make_png(tmp_path / "wide.png", width=1920, height=1080)
    document, _ = build_document(scenario, {"Scenario::1.1": [png]})

    shape = document.inline_shapes[0]
    assert shape.width == content_width(document)
    # Height keeps the 16:9 aspect ratio.
    assert shape.height == pytest.approx(shape.width * 1080 / 1920, rel=0.01)


def test_caption_is_written_under_the_picture(scenario, tmp_path):
    png = make_png(tmp_path / "c.png")
    screenshots = {"Scenario::1.1": [{"path": str(png), "caption": "Hasil login"}]}
    document, _ = build_document(scenario, screenshots)

    assert "Hasil login" in _texts(document)
    runs = [r for p in document.paragraphs for r in p.runs if r.text == "Hasil login"]
    assert runs and runs[0].italic is True


def test_screenshot_accepts_str_path_dict_and_dataclass(scenario, tmp_path):
    png = make_png(tmp_path / "x.png")
    for value in (str(png), png, {"path": str(png)}, Screenshot(path=png)):
        document, reports = build_document(scenario, {"Scenario::1.1": [value]})
        assert _image_count(document) == 1
        assert reports[0].screenshot_count == 1


def test_unknown_uid_is_ignored(scenario, tmp_path):
    png = make_png(tmp_path / "y.png")
    document, reports = build_document(scenario, {"Module Lain::9.9": [png]})
    assert _image_count(document) == 0
    assert reports[0].screenshot_count == 0


def test_missing_image_file_becomes_a_note_and_is_reported(scenario, tmp_path):
    ghost = tmp_path / "hilang.png"
    document, reports = build_document(scenario, {"Scenario::1.1": [ghost]})

    assert _image_count(document) == 0
    assert any("hilang.png" in text for text in _texts(document))
    assert reports[0].missing_files == [str(ghost)]
    assert reports[0].is_complete is False


def test_corrupt_image_does_not_abort_the_export(scenario, tmp_path):
    broken = tmp_path / "broken.png"
    broken.write_bytes(b"not really a png")
    document, reports = build_document(scenario, {"Scenario::1.1": [broken]})

    assert _image_count(document) == 0
    assert any("tidak bisa dibaca" in text for text in _texts(document))
    assert reports[0].missing_files == [str(broken)]


# --------------------------------------------------------------------------- #
# Titles and module selection
# --------------------------------------------------------------------------- #

def test_module_title_defaults_to_uppercase_sheet_name(tmp_path):
    path = simple_workbook(
        tmp_path / "amf.xlsx", sheet_name="New Validation AMF Segmentation"
    )
    document, _ = build_document(parse_workbook(path, config=CONFIG))
    assert _texts(document, "Heading 1") == ["NEW VALIDATION AMF SEGMENTATION"]


def test_project_name_metadata_becomes_the_module_title(tmp_path):
    from tests.fixtures import DEFAULT_METADATA

    path = simple_workbook(
        tmp_path / "meta.xlsx",
        sheet_name="Validation AMF Segmentation",
        header_row=9,
        metadata=DEFAULT_METADATA,          # Project Name = 'New Validation AMF'
        metadata_start_row=2,
    )
    scenario_file = parse_workbook(path, config=CONFIG)

    document, reports = build_document(scenario_file)
    assert _texts(document, "Heading 1") == ["NEW VALIDATION AMF"]
    assert reports[0].title == "NEW VALIDATION AMF"

    # The file name keeps the original capitalisation.
    result = export_per_module(scenario_file, output_dir=tmp_path / "out")[0]
    assert result.path.name == "Result Test UAT New Validation AMF.docx"


def test_module_name_metadata_wins_over_project_name(tmp_path):
    path = simple_workbook(
        tmp_path / "modname.xlsx",
        header_row=9,
        metadata=[("Project Name", "Proyek Besar"), ("Module", "AMF Segmentation")],
        metadata_start_row=2,
    )
    document, _ = build_document(parse_workbook(path, config=CONFIG))
    assert _texts(document, "Heading 1") == ["AMF SEGMENTATION"]


def test_sheet_name_is_used_when_there_is_no_metadata(scenario):
    document, _ = build_document(scenario)
    assert _texts(document, "Heading 1") == ["SCENARIO"]
    assert resolve_naming(scenario)["Scenario"].source == "sheet"


def test_workbook_project_name_only_titles_a_single_module_document(tmp_path):
    from tests.fixtures import DEFAULT_METADATA

    def builder(wb):
        # Only the first sheet carries the metadata block.
        write_scenario_sheet(
            wb.create_sheet("Module A"),
            header_row=9,
            metadata=DEFAULT_METADATA,
            metadata_start_row=2,
        )
        write_scenario_sheet(wb.create_sheet("Module B"))

    scenario_file = parse_workbook(build_workbook(tmp_path / "two.xlsx", builder), config=CONFIG)

    # Both modules in one document: the single project name cannot name both.
    document, _ = build_document(scenario_file)
    assert _texts(document, "Heading 1") == ["NEW VALIDATION AMF", "MODULE B"]

    # Module B on its own still falls back to its sheet name, not to the
    # project name inherited from Module A.
    document, _ = build_document(scenario_file, modules=["Module B"])
    assert _texts(document, "Heading 1") == ["MODULE B"]

    # Module A on its own uses its own metadata.
    document, _ = build_document(scenario_file, modules=["Module A"])
    assert _texts(document, "Heading 1") == ["NEW VALIDATION AMF"]


def test_modules_sharing_a_project_name_fall_back_to_sheet_names(tmp_path):
    shared = [("Project Name", "Proyek Yang Sama")]

    def builder(wb):
        for name in ("Module A", "Module B"):
            write_scenario_sheet(
                wb.create_sheet(name), header_row=9, metadata=shared, metadata_start_row=2
            )

    scenario_file = parse_workbook(build_workbook(tmp_path / "same.xlsx", builder), config=CONFIG)

    document, _ = build_document(scenario_file)
    assert _texts(document, "Heading 1") == ["MODULE A", "MODULE B"]

    # Exporting one file per module must not rename them either.
    results = export_per_module(scenario_file, output_dir=tmp_path / "out")
    assert [r.path.name for r in results] == [
        "Result Test UAT Module A.docx",
        "Result Test UAT Module B.docx",
    ]


def test_module_title_can_be_overridden(scenario):
    document, reports = build_document(scenario, titles={"Scenario": "JUDUL PILIHAN SAYA"})
    assert _texts(document, "Heading 1") == ["JUDUL PILIHAN SAYA"]
    assert reports[0].title == "JUDUL PILIHAN SAYA"


def test_override_is_used_verbatim_and_beats_metadata(tmp_path):
    from tests.fixtures import DEFAULT_METADATA

    path = simple_workbook(
        tmp_path / "ovr.xlsx", header_row=9, metadata=DEFAULT_METADATA, metadata_start_row=2
    )
    scenario_file = parse_workbook(path, config=CONFIG)
    titles = {"Scenario": "Judul Campur Huruf"}

    document, _ = build_document(scenario_file, titles=titles)
    # Not upper-cased: the user typed it exactly like this.
    assert _texts(document, "Heading 1") == ["Judul Campur Huruf"]

    result = export_per_module(scenario_file, output_dir=tmp_path / "o", titles=titles)[0]
    assert result.path.name == "Result Test UAT Judul Campur Huruf.docx"


def test_combined_document_contains_every_module(two_modules):
    document, reports = build_document(two_modules)
    assert _texts(document, "Heading 1") == ["AMF SEGMENTATION", "BILLING ADJUSTMENT"]
    assert len(reports) == 2
    assert len(_texts(document, "Heading 2")) == 4


def test_module_subset_selection(two_modules):
    document, reports = build_document(two_modules, modules=["Billing Adjustment"])
    assert _texts(document, "Heading 1") == ["BILLING ADJUSTMENT"]
    assert [r.module for r in reports] == ["Billing Adjustment"]


def test_unknown_module_raises(two_modules):
    with pytest.raises(ValueError, match="tidak ada"):
        build_document(two_modules, modules=["Tidak Ada"])


def test_page_break_between_modules(two_modules):
    document, _ = build_document(two_modules, page_break_between_modules=True)
    assert "w:br" in document.element.body.xml

    document, _ = build_document(two_modules, page_break_between_modules=False)
    assert 'type="page"' not in document.element.body.xml


# --------------------------------------------------------------------------- #
# Export to files
# --------------------------------------------------------------------------- #

def test_export_per_module_file_names(two_modules, tmp_path):
    out = tmp_path / "output"
    results = export_per_module(two_modules, output_dir=out)

    assert [r.path.name for r in results] == [
        "Result Test UAT AMF Segmentation.docx",
        "Result Test UAT Billing Adjustment.docx",
    ]
    for result in results:
        assert result.path.is_file()
        assert len(docx.Document(str(result.path)).paragraphs) > 0


def test_export_combined_uses_project_name(two_modules, tmp_path):
    out = tmp_path / "output"
    result = export_combined(two_modules, output_dir=out, project_name="AMF Segmentation")

    assert result.path.name == "Result Test UAT AMF Segmentation.docx"
    assert result.path.is_file()
    assert result.test_case_count == 4
    assert len(result.modules) == 2


def test_export_combined_falls_back_to_metadata_then_filename(tmp_path):
    from tests.fixtures import DEFAULT_METADATA

    path = simple_workbook(
        tmp_path / "meta.xlsx", header_row=9, metadata=DEFAULT_METADATA, metadata_start_row=2
    )
    scenario_file = parse_workbook(path, config=CONFIG)
    result = export_combined(scenario_file, output_dir=tmp_path / "out")
    assert result.path.name == "Result Test UAT New Validation AMF.docx"

    plain = parse_workbook(simple_workbook(tmp_path / "plain.xlsx"), config=CONFIG)
    result = export_combined(plain, output_dir=tmp_path / "out")
    assert result.path.name == "Result Test UAT plain.docx"


def test_export_dispatch(two_modules, tmp_path):
    assert len(export(two_modules, output_dir=tmp_path / "a")) == 2
    assert len(export(two_modules, output_dir=tmp_path / "b", combined=True)) == 1


def test_export_creates_the_output_folder(scenario, tmp_path):
    out = tmp_path / "deep" / "output"
    assert not out.exists()
    export_per_module(scenario, output_dir=out)
    assert out.is_dir()


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("AMF Segmentation", "AMF Segmentation"),
        ("Billing/Adjustment", "Billing Adjustment"),
        ('Modul: "A" <1>', "Modul A 1"),
        ("   ", "Output"),
    ],
)
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def test_default_output_name():
    assert default_output_name("AMF") == "Result Test UAT AMF.docx"


# --------------------------------------------------------------------------- #
# Template support
# --------------------------------------------------------------------------- #

@pytest.fixture
def template(tmp_path):
    """A user template: custom margins, a footer, and existing content."""
    doc = docx.Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    doc.sections[0].footer.paragraphs[0].text = "PT Contoh - Internal"
    doc.styles["Normal"].font.name = "Verdana"
    doc.add_heading("DOKUMEN LAMA YANG HARUS HILANG", 1)
    doc.add_paragraph("Isi lama yang harus hilang.")
    path = tmp_path / "template.docx"
    doc.save(str(path))
    return path


def test_template_keeps_style_and_footer_but_drops_its_content(scenario, template):
    document, _ = build_document(scenario, template=template)

    texts = _texts(document)
    assert "Isi lama yang harus hilang." not in texts
    assert "DOKUMEN LAMA YANG HARUS HILANG" not in texts
    assert _texts(document, "Heading 1") == ["SCENARIO"]

    assert document.styles["Normal"].font.name == "Verdana"
    assert document.sections[0].footer.paragraphs[0].text == "PT Contoh - Internal"
    assert document.sections[0].left_margin == Inches(0.5)


def test_template_margins_drive_the_screenshot_width(scenario, template, tmp_path):
    png = make_png(tmp_path / "t.png", width=1920, height=1080)
    document, _ = build_document(scenario, {"Scenario::1.1": [png]}, template=template)

    expected = (
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    )
    assert document.inline_shapes[0].width == expected


def test_template_file_is_not_modified(scenario, template, tmp_path):
    import hashlib

    before = hashlib.sha256(template.read_bytes()).hexdigest()
    export_per_module(scenario, output_dir=tmp_path / "out", template=template)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_template_images_do_not_leak_into_the_output(scenario, tmp_path):
    """A finished result document used as a template must not carry its
    old screenshots into the new file."""
    import zipfile

    doc = docx.Document()
    doc.add_heading("DOKUMEN LAMA", 1)
    doc.add_picture(str(make_png(tmp_path / "old.png", 1920, 1080)))
    heavy = tmp_path / "heavy_template.docx"
    doc.save(str(heavy))
    with zipfile.ZipFile(heavy) as zf:
        assert [n for n in zf.namelist() if "/media/" in n]

    out = tmp_path / "out"
    export_per_module(scenario, output_dir=out, template=heavy)
    produced = out / "Result Test UAT Scenario.docx"

    with zipfile.ZipFile(produced) as zf:
        assert [n for n in zf.namelist() if "/media/" in n] == []
    assert len(docx.Document(str(produced)).inline_shapes) == 0
    assert produced.stat().st_size < heavy.stat().st_size


def test_header_image_from_the_template_is_kept(scenario, tmp_path):
    """A logo lives in the header part, so it must survive the cleanup."""
    import zipfile

    doc = docx.Document()
    header = doc.sections[0].header
    header.paragraphs[0].add_run().add_picture(
        str(make_png(tmp_path / "logo.png", 200, 60)), width=Inches(1)
    )
    path = tmp_path / "logo_template.docx"
    doc.save(str(path))

    out = tmp_path / "out_logo"
    export_per_module(scenario, output_dir=out, template=path)
    produced = out / "Result Test UAT Scenario.docx"

    with zipfile.ZipFile(produced) as zf:
        assert [n for n in zf.namelist() if "/media/" in n]
    result = docx.Document(str(produced))
    assert "graphicData" in result.sections[0].header.paragraphs[0]._p.xml


def test_missing_template_raises(scenario, tmp_path):
    with pytest.raises(FileNotFoundError):
        build_document(scenario, template=tmp_path / "nope.docx")


def test_template_without_heading_styles_falls_back_to_bold(scenario, tmp_path):
    doc = docx.Document()
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].delete()
    path = tmp_path / "nostyles.docx"
    doc.save(str(path))

    document, _ = build_document(scenario, template=path)
    bold_runs = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "SCENARIO" in bold_runs
    assert any(text.startswith("1.1\t") for text in bold_runs)


# --------------------------------------------------------------------------- #
# Pre-export summary
# --------------------------------------------------------------------------- #

def test_summarize_without_writing_anything(two_modules, tmp_path):
    png = make_png(tmp_path / "s.png")
    reports = summarize(two_modules, {"AMF Segmentation::1.1": [png]})

    assert [r.module for r in reports] == ["AMF Segmentation", "Billing Adjustment"]
    assert reports[0].test_case_count == 2
    assert reports[0].screenshot_count == 1
    assert reports[0].without_screenshot == ["1.2  Validasi File Hasil Onboarding"]
    assert len(reports[1].without_screenshot) == 2
    assert list(tmp_path.glob("*.docx")) == []


def test_summarize_matches_the_built_document(two_modules, tmp_path):
    png = make_png(tmp_path / "m.png")
    screenshots = {"AMF Segmentation::1.1": [png, png]}

    pre = summarize(two_modules, screenshots)
    _document, post = build_document(two_modules, screenshots)

    assert [r.to_dict() for r in pre] == [r.to_dict() for r in post]


def test_summarize_reports_missing_files(scenario, tmp_path):
    reports = summarize(scenario, {"Scenario::1.1": [tmp_path / "ghost.png"]})
    assert reports[0].missing_files == [str(tmp_path / "ghost.png")]


def test_duplicate_test_case_id_keeps_original_heading_and_own_screenshot(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Login benar", "test_step": "a"},
        {"test_case_id": "1.2", "test_case": "Logout", "test_step": "b"},
        {"test_case_id": "1.1", "test_case": "Login salah", "test_step": "c"},
    ]
    scenario = parse_workbook(simple_workbook(tmp_path / "dup.xlsx", rows=rows), config=CONFIG)
    shot = make_png(tmp_path / "s.png")

    document, _reports = build_document(
        scenario, {"Scenario::1.1~2": [Screenshot(path=shot)]}
    )

    assert _texts(document, "Heading 2") == [
        "1.1\tLogin benar", "1.2\tLogout", "1.1\tLogin salah",
    ]
    assert _image_count(document) == 1
