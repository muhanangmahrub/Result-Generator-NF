"""Unit tests for the Excel scenario parser.

Covered scenarios required by the spec: shuffled column order, header row at
different positions, multi-sheet workbooks mixed with non-scenario sheets, and
Test Case IDs duplicated across sheets.
"""

from __future__ import annotations

import datetime as dt
import hashlib

import pytest
from openpyxl import Workbook, load_workbook

from rtg.config import load_config
from rtg.errors import (
    CorruptWorkbookError,
    MissingRequiredColumnsError,
    NoScenarioSheetError,
    ScenarioFileNotFoundError,
    SheetNotFoundError,
    UnsupportedFileFormatError,
)
from rtg.excel_parser import (
    SheetOverride,
    inspect_workbook,
    match_columns,
    parse_workbook,
)
from rtg.normalize import normalize_header
from tests.fixtures import (
    DEFAULT_METADATA,
    FULL_FIELD_ORDER,
    SAMPLE_ROWS,
    build_workbook,
    simple_workbook,
    write_noise_sheet,
    write_scenario_sheet,
)

CONFIG = load_config()


def _map(headers):
    """Convenience: map a list of raw header labels to fields."""
    return match_columns([normalize_header(h) for h in headers], CONFIG)


# --------------------------------------------------------------------------- #
# Column matching
# --------------------------------------------------------------------------- #

def test_exact_match_wins_over_prefix_match():
    # "Test Case" must not steal the "Test Case ID" column, in either order.
    mapping = _map(["Test Case ID", "Test Case", "Test Step"])
    assert mapping["test_case_id"] == 1
    assert mapping["test_case"] == 2
    assert mapping["test_step"] == 3

    mapping = _map(["Test Case", "Test Case ID", "Test Step"])
    assert mapping["test_case_id"] == 2
    assert mapping["test_case"] == 1


def test_match_handles_decorated_headers():
    mapping = _map(["Test Case ID / No", "Test Case Name", "Test Step", "Pass/Fail (P/F)"])
    assert mapping["test_case_id"] == 1
    assert mapping["test_case"] == 2
    assert mapping["pass_fail"] == 4


def test_match_ignores_unknown_headers():
    mapping = _map(["Nomor", "Test Case ID", "Test Case", "Test Step", "Lampiran"])
    assert mapping["test_case_id"] == 2
    assert "test_step_id" not in mapping
    assert set(mapping.values()).isdisjoint({1, 5})


def test_each_column_used_once():
    mapping = _map(["Test Case ID", "Test Case", "Test Step"])
    assert len(set(mapping.values())) == len(mapping)


# --------------------------------------------------------------------------- #
# Sheet detection
# --------------------------------------------------------------------------- #

def test_detects_scenario_sheets_regardless_of_name(tmp_path):
    def builder(wb: Workbook) -> None:
        write_noise_sheet(wb.create_sheet("Main"), "Main")
        write_noise_sheet(wb.create_sheet("Changes Logs"), "Changes Logs")
        write_scenario_sheet(wb.create_sheet("AMF Segmentation"))
        write_noise_sheet(wb.create_sheet("Question"), "Question")
        write_scenario_sheet(wb.create_sheet("Partner Onboarding"))
        write_noise_sheet(wb.create_sheet("Summary Scenario"), "Summary Scenario")
        write_noise_sheet(wb.create_sheet("ZIP Group"), "ZIP Group")

    path = build_workbook(tmp_path / "multi.xlsx", builder)
    inspection = inspect_workbook(path, config=CONFIG)

    assert inspection.scenario_sheet_names == ["AMF Segmentation", "Partner Onboarding"]
    assert len(inspection.sheets) == 7
    # A change log has Date/Description columns but no mandatory column, so it
    # must not even be offered for manual mapping.
    assert inspection.sheet("Changes Logs").is_mapping_candidate is False
    assert inspection.sheet("Main").is_scenario is False
    # "Total Scenario"/"Total Test Case" look like scenario columns but the
    # sheet still has fewer than two mandatory columns.
    assert inspection.sheet("Summary Scenario").is_mapping_candidate is False


def test_parse_all_scenario_sheets_by_default(tmp_path):
    def builder(wb: Workbook) -> None:
        write_noise_sheet(wb.create_sheet("Main"), "Main")
        write_scenario_sheet(wb.create_sheet("Module A"))
        write_scenario_sheet(wb.create_sheet("Module B"))

    path = build_workbook(tmp_path / "default.xlsx", builder)
    result = parse_workbook(path, config=CONFIG)

    assert result.module_names == ["Module A", "Module B"]
    assert len(result.test_cases) == 4


def test_selecting_a_subset_of_sheets(tmp_path):
    def builder(wb: Workbook) -> None:
        write_scenario_sheet(wb.create_sheet("Module A"))
        write_scenario_sheet(wb.create_sheet("Module B"))

    path = build_workbook(tmp_path / "subset.xlsx", builder)
    result = parse_workbook(path, ["Module B"], config=CONFIG)
    assert result.module_names == ["Module B"]


def test_unknown_sheet_name_raises(tmp_path):
    path = simple_workbook(tmp_path / "one.xlsx")
    with pytest.raises(SheetNotFoundError) as excinfo:
        parse_workbook(path, ["Tidak Ada"], config=CONFIG)
    assert "Tidak Ada" in excinfo.value.message
    assert "Scenario" in excinfo.value.message


def test_workbook_without_scenario_sheet_raises(tmp_path):
    def builder(wb: Workbook) -> None:
        write_noise_sheet(wb.create_sheet("Main"), "Main")
        write_noise_sheet(wb.create_sheet("ZIP Group"), "ZIP Group")

    path = build_workbook(tmp_path / "noscenario.xlsx", builder)
    with pytest.raises(NoScenarioSheetError) as excinfo:
        parse_workbook(path, config=CONFIG)
    assert "Tidak ada sheet skenario" in excinfo.value.message
    assert {s.name for s in excinfo.value.sheets} == {"Main", "ZIP Group"}


# --------------------------------------------------------------------------- #
# Shuffled columns / moved header row
# --------------------------------------------------------------------------- #

def test_shuffled_column_order(tmp_path):
    shuffled = (
        "notes",
        "pass_fail",
        "test_step",
        "test_case",
        "expected_results",
        "test_case_id",
        "test_step_id",
        "test_scenario",
        "scenario_id",
    )
    path = simple_workbook(tmp_path / "shuffled.xlsx", fields=shuffled)
    result = parse_workbook(path, config=CONFIG)
    module = result.modules[0]

    assert module.column_map["test_case_id"] == 6
    assert module.column_map["test_case"] == 4
    assert module.column_map["test_step"] == 3

    cases = module.test_cases
    assert [c.test_case_id for c in cases] == ["1.1", "1.2"]
    assert cases[0].test_case == "Move Partner Onboarding File via Sharefolder"
    assert [s.test_step for s in cases[0].steps] == [
        "Login ke server sharefolder",
        "Copy file ke folder tujuan",
    ]


@pytest.mark.parametrize("header_row", [1, 2, 9, 17])
def test_header_row_detected_at_any_position(tmp_path, header_row):
    path = simple_workbook(
        tmp_path / f"header{header_row}.xlsx",
        header_row=header_row,
        metadata=DEFAULT_METADATA if header_row > 7 else None,
        metadata_start_row=2,
    )
    result = parse_workbook(path, config=CONFIG)
    module = result.modules[0]
    assert module.header_row == header_row
    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2"]


def test_table_not_starting_in_column_a(tmp_path):
    path = simple_workbook(tmp_path / "offset.xlsx", header_row=5, start_col=3)
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert module.column_map["scenario_id"] == 3
    assert len(module.test_cases) == 2


def test_header_row_beyond_scan_window_is_not_detected(tmp_path):
    path = simple_workbook(tmp_path / "deep.xlsx", header_row=40)
    with pytest.raises(NoScenarioSheetError):
        parse_workbook(path, config=CONFIG)


# --------------------------------------------------------------------------- #
# Metadata
# --------------------------------------------------------------------------- #

def test_metadata_read_from_above_header(tmp_path):
    path = simple_workbook(
        tmp_path / "meta.xlsx",
        header_row=9,
        metadata=DEFAULT_METADATA,
        metadata_start_row=2,
    )
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert module.metadata["project_name"] == "New Validation AMF"
    assert module.metadata["version"] == "1.0"
    assert module.metadata["environment_testing"] == "UAT"
    assert module.metadata["type_testing"] == "Functional"
    assert module.metadata["created_by"] == "QA Engineer"


def test_metadata_value_can_be_far_to_the_right(tmp_path):
    path = simple_workbook(
        tmp_path / "meta_gap.xlsx",
        header_row=6,
        metadata=[("Project Name", "AMF")],
        metadata_start_row=2,
        metadata_label_col=1,
        metadata_value_col=4,
    )
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert module.metadata["project_name"] == "AMF"


def test_metadata_falls_back_to_another_sheet(tmp_path):
    def builder(wb: Workbook) -> None:
        write_scenario_sheet(
            wb.create_sheet("Module A"),
            header_row=9,
            metadata=DEFAULT_METADATA,
            metadata_start_row=2,
        )
        write_scenario_sheet(wb.create_sheet("Module B"))

    path = build_workbook(tmp_path / "metafallback.xlsx", builder)
    result = parse_workbook(path, config=CONFIG)

    assert result.module("Module B").metadata == {}
    assert result.merged_metadata()["project_name"] == "New Validation AMF"


def test_metadata_absent_is_not_an_error(tmp_path):
    path = simple_workbook(tmp_path / "nometa.xlsx")
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert module.metadata == {}


# --------------------------------------------------------------------------- #
# Fill-down
# --------------------------------------------------------------------------- #

def test_group_columns_are_filled_down(tmp_path):
    path = simple_workbook(tmp_path / "filldown.xlsx")
    module = parse_workbook(path, config=CONFIG).modules[0]
    case = module.test_cases[0]

    assert len(case.steps) == 2
    assert case.scenario_id == "SC-01"
    assert case.test_scenario == "Partner Onboarding"
    # Second test case inherits the scenario but keeps its own name.
    second = module.test_cases[1]
    assert second.scenario_id == "SC-01"
    assert second.test_case == "Validasi File Hasil Onboarding"


def test_new_test_case_id_does_not_inherit_previous_name(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Case Satu", "test_step": "step a"},
        {"test_step": "step b"},
        # New ID, name intentionally blank: it must not inherit "Case Satu".
        {"test_case_id": "1.2", "test_step": "step c"},
    ]
    path = simple_workbook(tmp_path / "reset.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2"]
    assert module.test_cases[0].test_case == "Case Satu"
    assert module.test_cases[1].test_case == ""


def test_merged_cells_behave_like_filled_cells(tmp_path):
    def builder(wb: Workbook) -> None:
        ws = wb.create_sheet("Merged")
        for col, label in enumerate(["Test Case ID", "Test Case", "Test Step"], start=1):
            ws.cell(row=1, column=col, value=label)
        ws.cell(row=2, column=1, value="1.1")
        ws.cell(row=2, column=2, value="Case Satu")
        ws.cell(row=2, column=3, value="step a")
        ws.cell(row=3, column=3, value="step b")
        ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
        ws.merge_cells(start_row=2, start_column=2, end_row=3, end_column=2)

    path = build_workbook(tmp_path / "merged.xlsx", builder)
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert len(module.test_cases) == 1
    assert len(module.test_cases[0].steps) == 2


# --------------------------------------------------------------------------- #
# Duplicate Test Case IDs
# --------------------------------------------------------------------------- #

def test_duplicate_test_case_id_across_sheets_stays_unique(tmp_path):
    def builder(wb: Workbook) -> None:
        write_scenario_sheet(wb.create_sheet("AMF Segmentation"))
        write_scenario_sheet(wb.create_sheet("Partner Onboarding"))

    path = build_workbook(tmp_path / "dup.xlsx", builder)
    result = parse_workbook(path, config=CONFIG)

    ids = [c.test_case_id for c in result.test_cases]
    uids = [c.uid for c in result.test_cases]
    assert ids == ["1.1", "1.2", "1.1", "1.2"]
    assert uids == [
        "AMF Segmentation::1.1",
        "AMF Segmentation::1.2",
        "Partner Onboarding::1.1",
        "Partner Onboarding::1.2",
    ]
    assert len(set(uids)) == 4
    assert result.test_case("Partner Onboarding::1.1").module == "Partner Onboarding"


def test_duplicate_test_case_id_within_a_sheet_is_kept_separate_with_warning(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Case Satu", "test_step": "step a"},
        {"test_case_id": "1.2", "test_case": "Case Dua", "test_step": "step b"},
        {"test_case_id": "1.1", "test_step": "step c"},
    ]
    path = simple_workbook(tmp_path / "dupsheet.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2", "1.1"]
    assert [c.uid for c in module.test_cases] == [
        "Scenario::1.1", "Scenario::1.2", "Scenario::1.1~2",
    ]
    assert [s.test_step for s in module.test_cases[0].steps] == ["step a"]
    assert [s.test_step for s in module.test_cases[2].steps] == ["step c"]
    warning = next(w for w in module.warnings if "muncul 2 kali" in w)
    assert "terpisah" in warning


def test_duplicate_id_warning_names_the_rows(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Case Satu", "test_step": "step a"},
        {"test_case_id": "1.2", "test_case": "Case Dua", "test_step": "step b"},
        {"test_case_id": "1.1", "test_case": "Case Satu Lagi", "test_step": "step c"},
    ]
    path = simple_workbook(tmp_path / "dup.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    first, third = module.test_cases[0], module.test_cases[2]
    warning = next(w for w in module.warnings if "muncul 2 kali" in w)
    assert f"{first.row}, {third.row}" in warning


def test_same_id_twice_in_a_row_with_different_names_is_two_cases(tmp_path):
    rows = [
        {"test_case_id": "2.1", "test_case": "Login benar", "test_step": "step a"},
        {"test_step": "step b"},
        {"test_case_id": "2.1", "test_case": "Login salah", "test_step": "step c"},
    ]
    path = simple_workbook(tmp_path / "samein.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    assert [c.test_case for c in module.test_cases] == ["Login benar", "Login salah"]
    assert [len(c.steps) for c in module.test_cases] == [2, 1]
    assert module.test_cases[1].uid == "Scenario::2.1~2"


def test_continuation_rows_are_not_split(tmp_path):
    rows = [
        {"test_case_id": "3.1", "test_case": "Case Tiga", "test_step": "step a"},
        {"test_case_id": "3.1", "test_case": "Case Tiga", "test_step": "step b"},
        {"test_case_id": "3.1", "test_step": "step c"},
    ]
    path = simple_workbook(tmp_path / "cont.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    assert len(module.test_cases) == 1
    assert len(module.test_cases[0].steps) == 3
    assert not any("muncul" in w for w in module.warnings)


def test_id_three_times_gets_occurrence_suffixes(tmp_path):
    rows = [
        {"test_case_id": "4.2", "test_case": "A", "test_step": "s1"},
        {"test_case_id": "5.1", "test_case": "B", "test_step": "s2"},
        {"test_case_id": "4.2", "test_case": "C", "test_step": "s3"},
        {"test_case_id": "5.2", "test_case": "D", "test_step": "s4"},
        {"test_case_id": "4.2", "test_case": "E", "test_step": "s5"},
    ]
    path = simple_workbook(tmp_path / "three.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    uids = [c.uid for c in module.test_cases if c.test_case_id == "4.2"]
    assert uids == ["Scenario::4.2", "Scenario::4.2~2", "Scenario::4.2~3"]
    assert len({c.uid for c in module.test_cases}) == len(module.test_cases)
    assert sum("'4.2'" in w for w in module.warnings) == 1
    assert any("muncul 3 kali" in w for w in module.warnings)


def test_duplicate_case_label_and_folder_are_distinguishable(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Sama", "test_step": "a"},
        {"test_case_id": "1.2", "test_case": "Lain", "test_step": "b"},
        {"test_case_id": "1.1", "test_case": "Sama", "test_step": "c"},
    ]
    path = simple_workbook(tmp_path / "label.xlsx", rows=rows)
    first, _, third = parse_workbook(path, config=CONFIG).modules[0].test_cases

    assert first.display_label != third.display_label
    assert "ID ganda #2" in third.display_label
    assert first.folder_name != third.folder_name
    assert first.heading == third.heading          # Word keeps the original ID


# --------------------------------------------------------------------------- #
# Optional columns, blank rows, repeated headers
# --------------------------------------------------------------------------- #

def test_only_mandatory_columns_present(tmp_path):
    path = simple_workbook(
        tmp_path / "minimal.xlsx",
        fields=("test_case_id", "test_case", "test_step"),
    )
    module = parse_workbook(path, config=CONFIG).modules[0]
    case = module.test_cases[0]
    assert case.scenario_id == ""
    assert case.steps[0].expected_results == ""
    assert len(module.scenarios) == 1
    assert module.scenarios[0].title == "(tanpa skenario)"


def test_optional_values_are_captured(tmp_path):
    rows = [{
        "test_case_id": "1.1",
        "test_case": "Case Satu",
        "test_step_id": "1",
        "test_step": "step a",
        "pre_condition": "User login",
        "test_data": "akun dummy",
        "expected_results": "muncul notifikasi",
        "actual_results": "sesuai",
        "pass_fail": "P",
        "date_executed": dt.datetime(2026, 10, 5),
        "test_executed_by": "Mahrub",
        "notes": "oke",
    }]
    path = simple_workbook(tmp_path / "optional.xlsx", rows=rows)
    step = parse_workbook(path, config=CONFIG).modules[0].test_cases[0].steps[0]

    assert step.test_step_id == "1"
    assert step.pre_condition == "User login"
    assert step.test_data == "akun dummy"
    assert step.expected_results == "muncul notifikasi"
    assert step.actual_results == "sesuai"
    assert step.pass_fail == "P"
    assert step.date_executed == "2026-10-05"
    assert step.test_executed_by == "Mahrub"
    assert step.notes == "oke"


def test_reading_stops_after_consecutive_blank_rows(tmp_path):
    rows = list(SAMPLE_ROWS) + [{}] * 12 + [
        {"test_case_id": "9.9", "test_case": "Sisa Tabel Lain", "test_step": "jangan dibaca"}
    ]
    path = simple_workbook(tmp_path / "blanks.xlsx", rows=rows)
    result = parse_workbook(path, config=CONFIG)
    module = result.modules[0]

    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2"]
    assert any("dihentikan" in w for w in module.warnings)


def test_small_blank_gap_does_not_stop_reading(tmp_path):
    rows = list(SAMPLE_ROWS) + [{}] * 3 + [
        {"test_case_id": "1.3", "test_case": "Case Tiga", "test_step": "step x"}
    ]
    path = simple_workbook(tmp_path / "gap.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2", "1.3"]


def test_repeated_header_row_inside_data_is_skipped(tmp_path):
    rows = [
        {"test_case_id": "1.1", "test_case": "Case Satu", "test_step": "step a"},
        {"test_case_id": "Test Case ID", "test_case": "Test Case", "test_step": "Test Step"},
        {"test_case_id": "1.2", "test_case": "Case Dua", "test_step": "step b"},
    ]
    path = simple_workbook(tmp_path / "repeat.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]
    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2"]


def test_stray_row_without_step_is_reported(tmp_path):
    # Mirrors a real file: a Test Scenario name typed on its own row, with no
    # Test Case ID and no Test Step under it.
    rows = [
        {"scenario_id": "SC-01", "test_scenario": "Alpha", "test_case_id": "1.1",
         "test_case": "Case Satu", "test_step": "step a"},
        {"test_scenario": "Scenario Nyasar Tanpa Test Case"},
        {"scenario_id": "SC-02", "test_scenario": "Beta", "test_case_id": "2.1",
         "test_case": "Case Dua", "test_step": "step b"},
    ]
    path = simple_workbook(tmp_path / "stray.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    # The stray row neither adds a step nor corrupts the next scenario.
    assert module.step_count == 2
    assert [g.scenario_id for g in module.scenarios] == ["SC-01", "SC-02"]
    assert [c.test_case_id for c in module.test_cases] == ["1.1", "2.1"]
    assert any("diabaikan" in w and "3" in w for w in module.warnings)


def test_sheet_with_header_but_no_data_rows(tmp_path):
    path = simple_workbook(tmp_path / "empty.xlsx", rows=[])
    result = parse_workbook(path, config=CONFIG)
    module = result.modules[0]
    assert module.test_cases == []
    assert any("tidak punya baris data" in w for w in module.warnings)


# --------------------------------------------------------------------------- #
# Missing mandatory columns -> manual mapping
# --------------------------------------------------------------------------- #

def test_missing_mandatory_column_reports_diagnostics(tmp_path):
    fields = tuple(f for f in FULL_FIELD_ORDER if f != "test_case")
    path = simple_workbook(tmp_path / "missing.xlsx", fields=fields)

    inspection = inspect_workbook(path, config=CONFIG)
    sheet = inspection.sheet("Scenario")
    assert sheet.is_scenario is False
    assert sheet.missing_required == ["test_case"]
    assert sheet.is_mapping_candidate is True
    assert "Kolom wajib tidak ditemukan" in sheet.reason

    with pytest.raises(MissingRequiredColumnsError) as excinfo:
        parse_workbook(path, ["Scenario"], config=CONFIG)
    error = excinfo.value
    assert error.sheet_name == "Scenario"
    assert error.missing == ["test_case"]
    assert error.header_row == 1
    assert any("Test Case ID" == text for _col, text in error.headers)


def test_manual_override_maps_an_unrecognised_header(tmp_path):
    path = simple_workbook(
        tmp_path / "override.xlsx",
        fields=("test_case_id", "test_case", "test_step"),
        labels={"test_case": "Nama Pengujian"},
    )
    # Without the override the sheet is unusable.
    with pytest.raises(MissingRequiredColumnsError):
        parse_workbook(path, ["Scenario"], config=CONFIG)

    overrides = {"Scenario": {"columns": {"test_case": 2}}}
    result = parse_workbook(path, ["Scenario"], config=CONFIG, overrides=overrides)
    case = result.modules[0].test_cases[0]
    assert case.test_case == "Move Partner Onboarding File via Sharefolder"

    # The same override expressed as a dataclass (as held in memory).
    result2 = parse_workbook(
        path, ["Scenario"], config=CONFIG,
        overrides={"Scenario": SheetOverride(columns={"test_case": 2})},
    )
    assert result2.modules[0].test_cases[0].test_case == case.test_case


def test_override_can_force_the_header_row(tmp_path):
    def builder(wb: Workbook) -> None:
        ws = wb.create_sheet("Scenario")
        # A decoy header row that also looks like a scenario table.
        write_scenario_sheet(ws, header_row=3, fields=("test_case_id", "test_case", "test_step"))
        ws.cell(row=1, column=1, value="Test Case ID")
        ws.cell(row=1, column=2, value="Test Case")
        ws.cell(row=1, column=3, value="Test Step")

    path = build_workbook(tmp_path / "forcerow.xlsx", builder)
    overrides = {"Scenario": {"header_row": 3}}
    module = parse_workbook(path, ["Scenario"], config=CONFIG, overrides=overrides).modules[0]
    assert module.header_row == 3
    assert [c.test_case_id for c in module.test_cases] == ["1.1", "1.2"]


def test_sheet_without_mandatory_columns_is_skipped_when_not_selected(tmp_path):
    def builder(wb: Workbook) -> None:
        write_scenario_sheet(wb.create_sheet("Good"))
        write_scenario_sheet(
            wb.create_sheet("Broken"),
            fields=tuple(f for f in FULL_FIELD_ORDER if f != "test_case"),
        )

    path = build_workbook(tmp_path / "mixed.xlsx", builder)
    result = parse_workbook(path, config=CONFIG)
    assert result.module_names == ["Good"]


# --------------------------------------------------------------------------- #
# File level errors and safety
# --------------------------------------------------------------------------- #

def test_missing_file(tmp_path):
    with pytest.raises(ScenarioFileNotFoundError):
        parse_workbook(tmp_path / "nope.xlsx", config=CONFIG)


def test_legacy_xls_is_rejected_with_a_clear_message(tmp_path):
    legacy = tmp_path / "old.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(UnsupportedFileFormatError) as excinfo:
        parse_workbook(legacy, config=CONFIG)
    assert ".xlsx" in excinfo.value.message


def test_corrupt_workbook(tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_text("ini bukan file excel", encoding="utf-8")
    with pytest.raises(CorruptWorkbookError):
        parse_workbook(broken, config=CONFIG)


def test_open_in_excel_adds_a_warning(tmp_path):
    path = simple_workbook(tmp_path / "locked.xlsx")
    (tmp_path / "~$locked.xlsx").write_bytes(b"lock")
    result = parse_workbook(path, config=CONFIG)
    assert any("terbuka di Excel" in w for w in result.warnings)


def test_source_workbook_is_never_modified(tmp_path):
    path = simple_workbook(tmp_path / "readonly.xlsx")
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    before_mtime = path.stat().st_mtime_ns

    inspect_workbook(path, config=CONFIG)
    parse_workbook(path, config=CONFIG)

    assert hashlib.sha256(path.read_bytes()).hexdigest() == before_hash
    assert path.stat().st_mtime_ns == before_mtime
    # And the workbook is still readable afterwards.
    wb = load_workbook(path)
    assert wb.sheetnames == ["Scenario"]
    wb.close()


# --------------------------------------------------------------------------- #
# Model-level expectations used by the later stages
# --------------------------------------------------------------------------- #

def test_headings_and_slugs(tmp_path):
    path = simple_workbook(tmp_path / "heading.xlsx", sheet_name="New Validation AMF Segmentation")
    module = parse_workbook(path, config=CONFIG).modules[0]
    case = module.test_cases[0]

    assert module.title == "NEW VALIDATION AMF SEGMENTATION"
    assert module.slug == "new-validation-amf-segmentation"
    # The reference Word document separates ID and name with a tab.
    assert case.heading == "1.1\tMove Partner Onboarding File via Sharefolder"
    assert case.display_label == "1.1  Move Partner Onboarding File via Sharefolder"
    assert case.folder_name == "1.1_move-partner-onboarding-file-via-sharefolder"
    assert case.uid == "New Validation AMF Segmentation::1.1"


def test_scenario_grouping_and_counts(tmp_path):
    rows = [
        {"scenario_id": "SC-01", "test_scenario": "Alpha",
         "test_case_id": "1.1", "test_case": "A1", "test_step": "s1"},
        {"test_case_id": "1.2", "test_case": "A2", "test_step": "s2"},
        {"scenario_id": "SC-02", "test_scenario": "Beta",
         "test_case_id": "2.1", "test_case": "B1", "test_step": "s3"},
    ]
    path = simple_workbook(tmp_path / "groups.xlsx", rows=rows)
    module = parse_workbook(path, config=CONFIG).modules[0]

    assert [g.scenario_id for g in module.scenarios] == ["SC-01", "SC-02"]
    assert [len(g.test_cases) for g in module.scenarios] == [2, 1]
    assert module.step_count == 3
    assert module.scenarios[0].title == "SC-01  Alpha"


def test_to_dict_is_json_serialisable(tmp_path):
    import json

    path = simple_workbook(tmp_path / "json.xlsx")
    result = parse_workbook(path, config=CONFIG)
    data = json.loads(json.dumps(result.to_dict(), ensure_ascii=False))
    assert data["modules"][0]["scenarios"][0]["test_cases"][0]["test_case_id"] == "1.1"


def test_inspection_reports_unmapped_columns_and_row_estimate(tmp_path):
    path = simple_workbook(
        tmp_path / "unmapped.xlsx",
        fields=("test_case_id", "test_case", "test_step"),
        labels={},
    )
    # Add a column the aliases do not know about.
    wb = load_workbook(path)
    wb["Scenario"].cell(row=1, column=4, value="Lampiran Bukti")
    wb.save(path)
    wb.close()

    sheet = inspect_workbook(path, config=CONFIG).sheet("Scenario")
    assert sheet.is_scenario is True
    assert sheet.data_row_estimate == 3
    assert ("Lampiran Bukti" in [text for _c, text in sheet.unmapped_columns])
