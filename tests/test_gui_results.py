"""Offscreen tests for the results panel and the phase-2 export commands."""

from __future__ import annotations

import datetime as dt

import pytest

pytest.importorskip("PySide6")

from PySide6.QtWidgets import QApplication, QMessageBox          # noqa: E402

from rtg.config import load_config                               # noqa: E402
from rtg.gui import main_window as mw                            # noqa: E402
from rtg.gui.main_window import MainWindow                       # noqa: E402
from rtg.gui.results_panel import COLUMNS, BulkResultDialog, ResultsPanel  # noqa: E402
from rtg.project import Project                                  # noqa: E402
from tests.fixtures import simple_workbook                       # noqa: E402

CONFIG = load_config()

ROWS = [
    {"test_case_id": "1.1", "test_case": "Case Satu", "test_step_id": "1.1.1",
     "test_step": "buka aplikasi", "pass_fail": "P", "test_executed_by": "Anang"},
    {"test_step_id": "1.1.2", "test_step": "login"},
    {"test_case_id": "1.2", "test_case": "Case Dua", "test_step_id": "1.2.1",
     "test_step": "cek laporan"},
]

COLUMN_INDEX = {field: index for index, (field, _t, _e) in enumerate(COLUMNS)}


@pytest.fixture(scope="session")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def excel(tmp_path):
    return simple_workbook(tmp_path / "Scenario.xlsx", rows=ROWS)


@pytest.fixture
def project(tmp_path, excel):
    return Project.create(tmp_path / "proj", excel, config=CONFIG)


@pytest.fixture
def panel(qapp, project):
    widget = ResultsPanel()
    widget.set_test_case(project, "Scenario::1.1")
    yield widget
    widget.close()


# --------------------------------------------------------------------------- #
# Results panel
# --------------------------------------------------------------------------- #

def test_panel_lists_every_step(panel, project):
    assert panel.table.rowCount() == 2
    assert panel.table.item(0, COLUMN_INDEX["step"]).text() == "1.1.1"
    assert panel.table.item(0, COLUMN_INDEX["test_step"]).text() == "buka aplikasi"
    assert "Case Satu" in panel.header.text()
    assert "salinan" in panel.header.text()


def test_panel_prefills_from_the_workbook(panel):
    # The scenario file already had Pass/Fail and Executed by on step 1.
    assert panel.table.item(0, COLUMN_INDEX["pass_fail"]).text() == "P"
    assert panel.table.item(0, COLUMN_INDEX["test_executed_by"]).text() == "Anang"
    assert panel.table.item(1, COLUMN_INDEX["pass_fail"]).text() == ""


def test_step_columns_are_read_only(panel):
    from PySide6.QtCore import Qt

    for field in ("step", "test_step"):
        item = panel.table.item(0, COLUMN_INDEX[field])
        assert not (item.flags() & Qt.ItemFlag.ItemIsEditable)
    editable = panel.table.item(0, COLUMN_INDEX["actual_results"])
    assert editable.flags() & Qt.ItemFlag.ItemIsEditable


def test_editing_a_cell_is_stored_in_the_project(panel, project):
    panel.table.item(0, COLUMN_INDEX["actual_results"]).setText("sesuai ekspektasi")
    assert project.result_for("Scenario::1.1", "1.1.1")["actual_results"] == "sesuai ekspektasi"

    # ...and survives a reload of the project.
    reopened = Project.open(project.folder, config=CONFIG)
    assert reopened.result_for("Scenario::1.1", "1.1.1")["actual_results"] == "sesuai ekspektasi"


def test_clearing_a_cell_removes_the_stored_value(panel, project):
    panel.table.item(0, COLUMN_INDEX["pass_fail"]).setText("F")
    assert project.result_for("Scenario::1.1", "1.1.1")["pass_fail"] == "F"
    panel.table.item(0, COLUMN_INDEX["pass_fail"]).setText("")
    assert "pass_fail" not in project.result_for("Scenario::1.1", "1.1.1")


def test_header_counts_filled_steps(panel, project):
    # A value shown from the workbook is not a stored result yet.
    assert "0 sudah ada hasilnya" in panel.header.text()
    assert project.results_count() == 0

    panel.table.item(0, COLUMN_INDEX["pass_fail"]).setText("F")
    assert "1 sudah ada hasilnya" in panel.header.text()


def test_mark_all_pass(panel, project):
    panel._mark_all_pass()
    today = dt.date.today().isoformat()
    for key in ("1.1.1", "1.1.2"):
        stored = project.result_for("Scenario::1.1", key)
        assert stored["pass_fail"] == "P"
        assert stored["date_executed"] == today
    assert panel.table.item(1, COLUMN_INDEX["pass_fail"]).text() == "P"


def test_clear_all(panel, project):
    panel._mark_all_pass()
    assert project.steps_with_results("Scenario::1.1") == 2
    panel._clear_all()
    assert project.steps_with_results("Scenario::1.1") == 0
    # The workbook's own values are shown again.
    assert panel.table.item(0, COLUMN_INDEX["pass_fail"]).text() == "P"


def test_apply_to_all_only_touches_the_given_fields(panel, project):
    panel.table.item(0, COLUMN_INDEX["actual_results"]).setText("punya sendiri")
    panel._apply_to_all({"pass_fail": "P"})

    stored = project.result_for("Scenario::1.1", "1.1.1")
    assert stored["actual_results"] == "punya sendiri"
    assert stored["pass_fail"] == "P"


def test_panel_without_a_test_case(qapp, project):
    widget = ResultsPanel()
    try:
        widget.set_test_case(project, None)
        assert widget.table.rowCount() == 0
        assert not widget.btn_bulk.isEnabled()
    finally:
        widget.close()


# --------------------------------------------------------------------------- #
# Bulk dialog
# --------------------------------------------------------------------------- #

def test_bulk_dialog_returns_only_filled_fields(qapp):
    dialog = BulkResultDialog(executed_by="Mahrub")
    try:
        dialog.use_date.setChecked(False)
        assert dialog.values() == {"test_executed_by": "Mahrub"}

        dialog.pass_fail.setCurrentText("P")
        dialog.actual.setText("sesuai")
        dialog.use_date.setChecked(True)
        values = dialog.values()
        assert values["pass_fail"] == "P"
        assert values["actual_results"] == "sesuai"
        assert values["date_executed"] == dt.date.today().isoformat()
    finally:
        dialog.close()


def test_bulk_dialog_empty_gives_nothing(qapp):
    dialog = BulkResultDialog()
    try:
        dialog.use_date.setChecked(False)
        assert dialog.values() == {}
    finally:
        dialog.close()


# --------------------------------------------------------------------------- #
# Main window wiring
# --------------------------------------------------------------------------- #

@pytest.fixture
def window(qapp, project, monkeypatch):
    monkeypatch.setattr(MainWindow, "_start_hotkey", lambda self: None)
    win = MainWindow(project)
    yield win
    win.close()


def test_results_tab_follows_the_tree_selection(window, project):
    case_item = window.tree.topLevelItem(0).child(0).child(1)      # test case 1.2
    window.tree.setCurrentItem(case_item)
    assert window.results.table.rowCount() == 1
    assert window.results.table.item(0, COLUMN_INDEX["step"]).text() == "1.2.1"


def test_write_excel_without_results_explains_itself(window, project, monkeypatch):
    shown = []
    monkeypatch.setattr(
        QMessageBox, "information", staticmethod(lambda *a, **k: shown.append(a[2]))
    )
    window.write_result_excel()
    assert shown and "Belum ada Actual Results" in shown[0]


def test_write_excel_produces_a_copy(window, project, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    project.set_results_for_test_case("Scenario::1.1", {
        "pass_fail": "P", "test_executed_by": "Mahrub",
    })
    target = tmp_path / "Hasil.xlsx"
    monkeypatch.setattr(
        mw.QFileDialog, "getSaveFileName",
        staticmethod(lambda *a, **k: (str(target), "")),
    )
    monkeypatch.setattr(QMessageBox, "exec", lambda self: None)

    window.write_result_excel()

    assert target.is_file()
    module = project.scenario_file.modules[0]
    wb = load_workbook(target, data_only=True)
    try:
        row = project.test_case("Scenario::1.1").steps[0].row
        assert wb["Scenario"].cell(row=row, column=module.column_map["pass_fail"]).value == "P"
    finally:
        wb.close()


def test_write_excel_cancelled_writes_nothing(window, project, monkeypatch):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    monkeypatch.setattr(
        mw.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: ("", ""))
    )
    window.write_result_excel()
    assert list(project.results_dir.glob("*.xlsx")) == []


def test_export_pdf_without_documents_explains_itself(window, monkeypatch):
    shown = []
    monkeypatch.setattr(
        QMessageBox, "information", staticmethod(lambda *a, **k: shown.append(a[2]))
    )
    window.export_pdf()
    assert shown and "belum berisi file .docx" in shown[0]


def test_export_pdf_without_a_converter_explains_itself(window, project, monkeypatch):
    from rtg.pdf_export import ConverterInfo

    project.generate()                       # puts a .docx in output/
    monkeypatch.setattr(
        mw, "available_converters",
        lambda: [ConverterInfo("word", False, "tidak terpasang"),
                 ConverterInfo("libreoffice", False, "tidak ditemukan")],
    )
    shown = []
    monkeypatch.setattr(
        QMessageBox, "warning", staticmethod(lambda *a, **k: shown.append(a[2]))
    )
    window.export_pdf()
    assert shown and "butuh Microsoft Word" in shown[0]
    assert "tidak terpasang" in shown[0]


def test_orphan_results_reach_the_warnings_tab(window, project, excel):
    project.set_result("Scenario::1.2", "1.2.1", {"pass_fail": "P"})
    simple_workbook(excel, rows=ROWS[:2])
    project.reimport()
    window._refresh_all()

    assert "Hasil eksekusi tersimpan" in window.warnings.toPlainText()
    assert "Scenario::1.2" in window.warnings.toPlainText()
