"""Unit tests for the execution results stored in a project."""

from __future__ import annotations

import datetime as dt
import json

import pytest
from openpyxl import load_workbook

from rtg.config import load_config
from rtg.project import Project, ProjectError, step_key
from tests.fixtures import simple_workbook

CONFIG = load_config()

ROWS = [
    {"test_case_id": "1.1", "test_case": "Case Satu", "test_step_id": "1.1.1",
     "test_step": "buka aplikasi"},
    {"test_step_id": "1.1.2", "test_step": "login"},
    {"test_case_id": "1.2", "test_case": "Case Dua", "test_step_id": "1.2.1",
     "test_step": "cek laporan"},
]

ROWS_NO_STEP_ID = [
    {"test_case_id": "1.1", "test_case": "Case Satu", "test_step": "buka aplikasi"},
    {"test_step": "login"},
]


@pytest.fixture
def excel(tmp_path):
    return simple_workbook(tmp_path / "src" / "Scenario.xlsx", rows=ROWS)


@pytest.fixture
def project(tmp_path, excel):
    return Project.create(tmp_path / "proj", excel, config=CONFIG)


# --------------------------------------------------------------------------- #
# Step keys
# --------------------------------------------------------------------------- #

def test_step_key_prefers_the_test_step_id(project):
    case = project.test_case("Scenario::1.1")
    assert [step_key(s, i) for i, s in enumerate(case.steps, 1)] == ["1.1.1", "1.1.2"]


def test_step_key_falls_back_to_position(tmp_path):
    excel = simple_workbook(tmp_path / "nostepid.xlsx", rows=ROWS_NO_STEP_ID)
    project = Project.create(tmp_path / "p", excel, config=CONFIG)
    case = project.test_case("Scenario::1.1")
    assert [step_key(s, i) for i, s in enumerate(case.steps, 1)] == ["#1", "#2"]


# --------------------------------------------------------------------------- #
# Storing results
# --------------------------------------------------------------------------- #

def test_set_and_read_a_result(project):
    project.set_result("Scenario::1.1", "1.1.1", {
        "actual_results": "sesuai", "pass_fail": "P",
        "date_executed": "2026-10-27", "test_executed_by": "Mahrub",
    })
    stored = project.result_for("Scenario::1.1", "1.1.1")
    assert stored["actual_results"] == "sesuai"
    assert stored["pass_fail"] == "P"
    assert project.results_count() == 1
    assert project.steps_with_results("Scenario::1.1") == 1


def test_values_are_trimmed_and_blanks_clear_the_field(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "  P  ", "actual_results": "x"})
    assert project.result_for("Scenario::1.1", "1.1.1")["pass_fail"] == "P"

    project.set_result("Scenario::1.1", "1.1.1", {"actual_results": "   "})
    stored = project.result_for("Scenario::1.1", "1.1.1")
    assert "actual_results" not in stored
    assert stored["pass_fail"] == "P"          # untouched field survives


def test_clearing_every_field_drops_the_entry(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": ""})
    assert project.results_count() == 0
    assert "Scenario::1.1" not in project.results


def test_clear_result(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    project.clear_result("Scenario::1.1", "1.1.1")
    assert project.results_count() == 0


def test_unknown_field_is_rejected(project):
    with pytest.raises(ProjectError, match="tidak dikenali"):
        project.set_result("Scenario::1.1", "1.1.1", {"expected_results": "x"})


def test_fill_every_step_of_a_test_case(project):
    count = project.set_results_for_test_case("Scenario::1.1", {
        "pass_fail": "P", "test_executed_by": "Mahrub",
    })
    assert count == 2
    assert project.steps_with_results("Scenario::1.1") == 2
    assert project.result_for("Scenario::1.1", "1.1.2")["pass_fail"] == "P"


def test_fill_unknown_test_case(project):
    with pytest.raises(ProjectError, match="tidak ada"):
        project.set_results_for_test_case("Scenario::9.9", {"pass_fail": "P"})


def test_results_survive_save_and_open(project):
    project.set_result("Scenario::1.2", "1.2.1", {
        "actual_results": "laporan tampil", "pass_fail": "P",
    })
    reopened = Project.open(project.folder, config=CONFIG)
    assert reopened.result_for("Scenario::1.2", "1.2.1") == {
        "actual_results": "laporan tampil", "pass_fail": "P",
    }

    data = json.loads(project.project_file.read_text(encoding="utf-8"))
    assert data["results"]["Scenario::1.2"]["1.2.1"]["pass_fail"] == "P"


def test_unknown_fields_in_project_json_are_dropped_on_load(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    data = json.loads(project.project_file.read_text(encoding="utf-8"))
    data["results"]["Scenario::1.1"]["1.1.1"]["bukan_field"] = "x"
    project.project_file.write_text(json.dumps(data), encoding="utf-8")

    reopened = Project.open(project.folder, config=CONFIG)
    assert reopened.result_for("Scenario::1.1", "1.1.1") == {"pass_fail": "P"}


# --------------------------------------------------------------------------- #
# Re-import
# --------------------------------------------------------------------------- #

def test_results_follow_the_step_id_after_rows_move(project, excel):
    project.set_result("Scenario::1.2", "1.2.1", {"pass_fail": "P"})
    row_before = project.test_case("Scenario::1.2").steps[0].row

    # Push the table down by starting the header lower.
    simple_workbook(excel, rows=ROWS, header_row=6)
    project.reimport()

    row_after = project.test_case("Scenario::1.2").steps[0].row
    assert row_after != row_before
    assert project.result_for("Scenario::1.2", "1.2.1") == {"pass_fail": "P"}
    updates = project.result_updates()
    assert updates["Scenario"][0].row == row_after


def test_orphan_results_are_reported(project, excel):
    project.set_result("Scenario::1.2", "1.2.1", {"pass_fail": "P"})
    simple_workbook(excel, rows=ROWS[:2])          # test case 1.2 disappears
    project.reimport()

    assert project.orphan_results() == ["Scenario::1.2"]
    assert project.result_updates() == {}          # nothing to write
    # The data is kept, so it comes back with the test case.
    simple_workbook(excel, rows=ROWS)
    project.reimport()
    assert project.orphan_results() == []
    assert project.result_updates()["Scenario"]


# --------------------------------------------------------------------------- #
# Write-back
# --------------------------------------------------------------------------- #

def test_result_updates_shape(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    project.set_result("Scenario::1.2", "1.2.1", {"pass_fail": "F"})
    updates = project.result_updates()

    assert list(updates) == ["Scenario"]
    rows = {u.row: u.values for u in updates["Scenario"]}
    assert len(rows) == 2
    assert "P" in [v["pass_fail"] for v in rows.values()]


def test_write_result_workbook_default_location(project, excel):
    project.set_results_for_test_case("Scenario::1.1", {
        "pass_fail": "P", "date_executed": "2026-10-27", "test_executed_by": "Mahrub",
    })
    report = project.write_result_workbook()

    assert report.path.parent == project.results_dir
    assert report.path.name == "Scenario - Result.xlsx"
    assert report.cells_written == 6               # 2 steps x 3 fields
    assert report.path.is_file()

    module = project.scenario_file.modules[0]
    wb = load_workbook(report.path, data_only=True)
    try:
        sheet = wb["Scenario"]
        row = project.test_case("Scenario::1.1").steps[0].row
        assert sheet.cell(row=row, column=module.column_map["pass_fail"]).value == "P"
        assert sheet.cell(row=row, column=module.column_map["test_executed_by"]).value == "Mahrub"
    finally:
        wb.close()

    # The original is untouched.
    original = load_workbook(excel, data_only=True)
    try:
        assert original["Scenario"].cell(
            row=row, column=module.column_map["pass_fail"]
        ).value is None
    finally:
        original.close()


def test_write_result_workbook_to_a_chosen_path(project, tmp_path):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    destination = tmp_path / "pilihan" / "Hasil UAT.xlsx"
    report = project.write_result_workbook(destination)
    assert report.path == destination
    assert destination.is_file()


def test_write_result_workbook_without_any_result(project):
    with pytest.raises(ProjectError, match="Belum ada hasil"):
        project.write_result_workbook()


def test_write_result_workbook_rewrites_the_default_file(project):
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "P"})
    first = project.write_result_workbook()
    project.set_result("Scenario::1.1", "1.1.1", {"pass_fail": "F"})
    second = project.write_result_workbook()
    assert first.path == second.path               # overwritten, not refused

    module = project.scenario_file.modules[0]
    wb = load_workbook(second.path, data_only=True)
    try:
        row = project.test_case("Scenario::1.1").steps[0].row
        assert wb["Scenario"].cell(row=row, column=module.column_map["pass_fail"]).value == "F"
    finally:
        wb.close()


def test_results_are_independent_per_module(tmp_path):
    from tests.fixtures import build_workbook, write_scenario_sheet

    def builder(wb):
        write_scenario_sheet(wb.create_sheet("Module A"), rows=ROWS)
        write_scenario_sheet(wb.create_sheet("Module B"), rows=ROWS)

    excel = build_workbook(tmp_path / "two.xlsx", builder)
    project = Project.create(tmp_path / "p", excel, config=CONFIG)

    project.set_result("Module A::1.1", "1.1.1", {"pass_fail": "P"})
    project.set_result("Module B::1.1", "1.1.1", {"pass_fail": "F"})

    updates = project.result_updates()
    assert set(updates) == {"Module A", "Module B"}
    assert updates["Module A"][0].values["pass_fail"] == "P"
    assert updates["Module B"][0].values["pass_fail"] == "F"

    report = project.write_result_workbook()
    assert sorted(report.sheets) == ["Module A", "Module B"]
