"""Helpers that build synthetic scenario workbooks for the parser tests.

Everything is generated with openpyxl so the test suite needs no binary
fixtures, and each test can freely move the header row, shuffle the column
order, rename headers or drop optional columns.
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def make_png(path: Path, width: int = 320, height: int = 180,
             color: Tuple[int, int, int] = (40, 90, 160)) -> Path:
    """Write a valid solid-colour PNG, so the tests need no image library."""
    row = b"\x00" + bytes(color) * width
    raw = row * height

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)  # 8-bit truecolour
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path

#: Default header label used for each field.
DEFAULT_LABELS: Dict[str, str] = {
    "scenario_id": "Scenario ID",
    "test_scenario": "Test Scenario",
    "test_scenario_description": "Test Scenario Description",
    "test_case_id": "Test Case ID",
    "test_case": "Test Case",
    "test_step_id": "Test Step ID",
    "test_step": "Test Step",
    "positive_negative": "Positive/Negative (P/N)",
    "pre_condition": "Pre-Condition",
    "test_data": "Test Data",
    "expected_results": "Expected Results",
    "actual_results": "Actual Results",
    "pass_fail": "Pass/Fail (P/F)",
    "date_executed": "Date Executed",
    "test_executed_by": "Test Executed by",
    "notes": "Notes",
}

FULL_FIELD_ORDER: Tuple[str, ...] = (
    "scenario_id",
    "test_scenario",
    "test_case_id",
    "test_case",
    "test_step_id",
    "test_step",
    "pre_condition",
    "test_data",
    "expected_results",
    "actual_results",
    "pass_fail",
    "date_executed",
    "test_executed_by",
    "notes",
)

#: Two test cases, three steps, with group columns filled only on the first row
#: of each group - exactly how the real scenario files are written.
SAMPLE_ROWS: Tuple[Dict[str, Any], ...] = (
    {
        "scenario_id": "SC-01",
        "test_scenario": "Partner Onboarding",
        "test_case_id": "1.1",
        "test_case": "Move Partner Onboarding File via Sharefolder",
        "test_step_id": "1",
        "test_step": "Login ke server sharefolder",
        "expected_results": "Login berhasil",
        "pass_fail": "Pass",
    },
    {
        "test_step_id": "2",
        "test_step": "Copy file ke folder tujuan",
        "expected_results": "File tersalin",
        "pass_fail": "Pass",
    },
    {
        "test_case_id": "1.2",
        "test_case": "Validasi File Hasil Onboarding",
        "test_step_id": "1",
        "test_step": "Buka file hasil onboarding",
        "expected_results": "File terbuka dan format sesuai",
        "pass_fail": "Pass",
    },
)

DEFAULT_METADATA: Tuple[Tuple[str, str], ...] = (
    ("Project Name", "New Validation AMF"),
    ("Version", "1.0"),
    ("Environment Testing", "UAT"),
    ("Type Testing", "Functional"),
    ("Created By", "QA Engineer"),
    ("Created Date", "2026-10-01"),
)


def write_scenario_sheet(
    ws: Worksheet,
    *,
    fields: Sequence[str] = FULL_FIELD_ORDER,
    rows: Sequence[Dict[str, Any]] = SAMPLE_ROWS,
    header_row: int = 1,
    start_col: int = 1,
    labels: Optional[Dict[str, str]] = None,
    metadata: Optional[Iterable[Tuple[str, Any]]] = None,
    metadata_start_row: int = 1,
    metadata_label_col: int = 1,
    metadata_value_col: Optional[int] = None,
    title: Optional[str] = None,
) -> Worksheet:
    """Write a scenario table, with optional metadata block above the header."""
    label_map = dict(DEFAULT_LABELS)
    if labels:
        label_map.update(labels)

    if title:
        ws.cell(row=1, column=start_col, value=title)

    if metadata:
        value_col = metadata_value_col or (metadata_label_col + 1)
        for offset, (label, value) in enumerate(metadata):
            ws.cell(row=metadata_start_row + offset, column=metadata_label_col, value=label)
            ws.cell(row=metadata_start_row + offset, column=value_col, value=value)

    for offset, field in enumerate(fields):
        ws.cell(row=header_row, column=start_col + offset, value=label_map[field])

    for row_offset, data in enumerate(rows, start=1):
        for col_offset, field in enumerate(fields):
            value = data.get(field)
            if value is None or value == "":
                continue
            ws.cell(
                row=header_row + row_offset,
                column=start_col + col_offset,
                value=value,
            )
    return ws


def write_noise_sheet(ws: Worksheet, kind: str) -> Worksheet:
    """Write one of the non-scenario sheets found in the real files."""
    if kind == "Main":
        ws["A1"] = "Document Title"
        ws["B1"] = "Scenario Test UAT"
        ws["A2"] = "Project Name"
        ws["B2"] = "New Validation AMF"
        ws["A3"] = "Prepared By"
        ws["B3"] = "QA Engineer"
    elif kind == "Changes Logs":
        for col, label in enumerate(["No", "Date", "Description", "Author"], start=1):
            ws.cell(row=1, column=col, value=label)
        ws.append([1, "2026-10-01", "Initial version", "QA Engineer"])
    elif kind == "Question":
        ws["A1"] = "No"
        ws["B1"] = "Question"
        ws["C1"] = "Answer"
        ws.append([1, "Apakah data dummy tersedia?", "Ya"])
    elif kind == "Summary Scenario":
        ws["A1"] = "Module"
        ws["B1"] = "Total Scenario"
        ws["C1"] = "Total Test Case"
        ws.append(["AMF Segmentation", 1, 2])
    elif kind == "ZIP Group":
        ws["A1"] = "Group"
        ws["B1"] = "File Name"
        ws.append(["G1", "onboarding.zip"])
    else:  # pragma: no cover - guard for typos in tests
        raise ValueError(f"unknown noise sheet: {kind}")
    return ws


def build_workbook(path: Path, builder) -> Path:
    """Create a workbook at ``path``; ``builder(wb)`` adds the sheets."""
    wb = Workbook()
    wb.remove(wb.active)
    builder(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def simple_workbook(path: Path, **kwargs) -> Path:
    """A workbook with a single scenario sheet named 'Scenario'."""
    sheet_name = kwargs.pop("sheet_name", "Scenario")

    def builder(wb: Workbook) -> None:
        write_scenario_sheet(wb.create_sheet(sheet_name), **kwargs)

    return build_workbook(path, builder)
